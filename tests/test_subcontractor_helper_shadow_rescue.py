"""Phase 1.1 end-to-end + regression tests (Plan 01.1-05).

D-20 / Pitfall 6 closure: drives ``_subcontractor_rescue_price`` +
``group_source_rows`` + ``cleanup_untracked_sheet_attachments`` +
``_run_phase_1_1_hash_prune`` on synthetic Smartsheet payloads that
exercise all four upstream fixes (Bug A pre-acceptance rescue,
Bug B1 partitioning, Bug B2 cleanup whitelist, Bug C claim-history
attribution) AND the SUB-12 hash-history one-time prune.

Closes the structural weakness of
``tests/test_subcontractor_pricing.py::TestHelperShadowVariantFileIdentifier``
which is a mirror class — it bypasses both the upstream classifier
and the has_price gate, so it passes EVEN IF Bug A or Bug B1 is
broken in production (the failure mode that allowed Phase 1 to
ship with both bugs latent).

Per the Phase 1.1 Living Ledger entry rule (d): any plan that fixes
a row-flow bug — acceptance gate, ``group_source_rows``,
``generate_excel`` — MUST add at least one true end-to-end test
driving the full pipeline. Static mirror classes don't count.
"""

from __future__ import annotations

import datetime
import importlib
import inspect
import json
import os
import pathlib
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

_REPO_ROOT = Path(__file__).resolve().parent.parent
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

# Reuse fixture helpers from the existing billing_audit test file.
# Co-locating them in this Phase 1.1 file would duplicate the chained-
# mock builder; importing keeps the SoR single (see PATTERNS.md
# anchor 6 — "structural mirror" pattern).
from tests.test_billing_audit_shadow import (
    _reset_all,
    _ensure_smartsheet_mocked,
    _make_fake_supabase_client,
    _fake_rpc_response,
)

_ensure_smartsheet_mocked()
import generate_weekly_pdfs  # noqa: E402 — must come after mock injection
import pipeline.grouping as _grouping  # noqa: E402
from billing_audit.writer import ResolveOutcome  # noqa: E402


def _safe_reload_gwp():
    """Reload generate_weekly_pdfs.py with Sentry init suppressed.

    Mirrors the pattern in ``tests/test_performance_optimizations.py``
    — a dev shell with a real ``SENTRY_DSN`` would otherwise fire a
    live Sentry init on every reload. Phase 1.1 introduces two new
    module-level kill switches that the tests need to flip via env
    vars, so reload-with-suppression is the canonical setup pattern.
    """
    with mock.patch.dict(os.environ, {"SENTRY_DSN": ""}, clear=False), \
         mock.patch("sentry_sdk.init"):
        importlib.reload(generate_weekly_pdfs)


class TestEndToEndPipeline(unittest.TestCase):
    """D-20 / Pitfall 6: pipeline coverage for the four upstream fixes.

    Exercises SUB-08 (Bug A), SUB-09 (Bug B1), SUB-11 (Bug C) through
    real production code paths. The synthetic Smartsheet payload
    travels through ``_subcontractor_rescue_price`` (post-rescue
    price simulated for non-Bug-A tests) and ``group_source_rows``
    (variant emission + per-row attribution). Assertions are on
    observable outputs (group dict keys) NOT on internal helper
    return values.
    """

    _SUB_SHEET_ID = 8162920222379908
    _NON_SUB_SHEET_ID = 9999999999

    def setUp(self):
        _reset_all()
        # Snapshot module state so mutations don't leak across tests
        self._orig_enabled = generate_weekly_pdfs.SUBCONTRACTOR_RATE_VARIANTS_ENABLED
        self._orig_bug_a = generate_weekly_pdfs.SUBCONTRACTOR_RATE_RECALC_PREACCEPTANCE_ENABLED
        self._orig_bug_c = generate_weekly_pdfs.SUBCONTRACTOR_HELPER_CLAIM_ATTRIBUTION_ENABLED
        self._orig_sub_ids = set(generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS)
        self._orig_orig_ids = set(generate_weekly_pdfs._FOLDER_DISCOVERED_ORIG_IDS)
        self._orig_rates = dict(generate_weekly_pdfs._SUBCONTRACTOR_RATES)
        # Defaults for the test scenarios
        generate_weekly_pdfs.SUBCONTRACTOR_RATE_VARIANTS_ENABLED = True
        generate_weekly_pdfs.SUBCONTRACTOR_RATE_RECALC_PREACCEPTANCE_ENABLED = True
        generate_weekly_pdfs.SUBCONTRACTOR_HELPER_CLAIM_ATTRIBUTION_ENABLED = True
        generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS.clear()
        generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS.add(self._SUB_SHEET_ID)
        # Seed a rate so Bug A rescue can fire
        generate_weekly_pdfs._SUBCONTRACTOR_RATES['ANC-M'] = {
            'new_install_price': 75.0,
            'reduced_install_price': 50.0,
            'new_remove_price': 60.0,
            'reduced_remove_price': 45.0,
            'new_transfer_price': 80.0,
            'reduced_transfer_price': 55.0,
        }

    def tearDown(self):
        generate_weekly_pdfs.SUBCONTRACTOR_RATE_VARIANTS_ENABLED = self._orig_enabled
        generate_weekly_pdfs.SUBCONTRACTOR_RATE_RECALC_PREACCEPTANCE_ENABLED = self._orig_bug_a
        generate_weekly_pdfs.SUBCONTRACTOR_HELPER_CLAIM_ATTRIBUTION_ENABLED = self._orig_bug_c
        generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS.clear()
        generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS.update(self._orig_sub_ids)
        generate_weekly_pdfs._FOLDER_DISCOVERED_ORIG_IDS.clear()
        generate_weekly_pdfs._FOLDER_DISCOVERED_ORIG_IDS.update(self._orig_orig_ids)
        generate_weekly_pdfs._SUBCONTRACTOR_RATES.clear()
        generate_weekly_pdfs._SUBCONTRACTOR_RATES.update(self._orig_rates)
        _reset_all()

    def _make_synth_helper_row(
        self,
        wr='19236776',
        helper_foreman='ReplacementForeman',
        units_price='$100.00',
        snapshot='2026-04-19',
        row_id=12345,
        source_sheet_id=None,
    ):
        """Synthetic helper row that triggers Bug B1 partitioning + Bug C
        attribution. ``units_price`` defaults to ``'$100.00'`` (post-rescue
        state — Bug A's rescue is upstream of ``group_source_rows`` so the
        downstream tests simulate the post-rescue row directly per
        RESEARCH.md Pitfall 6: "the test can drive the rescue helper
        independently if needed").
        """
        return {
            '__row_id': row_id,
            'Work Request #': wr,
            'Weekly Reference Logged Date': '2026-04-19',
            'Snapshot Date': snapshot,
            'Units Completed?': True,
            'Foreman Helping?': helper_foreman,
            'Helping Foreman Completed Unit?': True,
            'Units Total Price': units_price,
            'CU': 'ANC-M',
            'Work Type': 'Inst',
            'Quantity': 2,
            '__effective_user': 'PrimaryForeman',
            '__assignment_method': 'FOREMAN_COLUMN',
            '__is_helper_row': True,
            '__helper_foreman': helper_foreman,
            '__helper_dept': '500',
            '__helper_job': 'JOB-A',
            '__is_vac_crew': False,
            '__source_sheet_id': source_sheet_id or self._SUB_SHEET_ID,
        }

    def _make_synth_non_helper_row(
        self,
        wr='19236776',
        units_price='$100.00',
        snapshot='2026-04-19',
        source_sheet_id=None,
    ):
        """Synthetic non-helper subcontractor row for Bug B1 partitioning."""
        return {
            'Work Request #': wr,
            'Weekly Reference Logged Date': '2026-04-19',
            'Snapshot Date': snapshot,
            'Units Completed?': True,
            'Units Total Price': units_price,
            'CU': 'ANC-M',
            'Work Type': 'Inst',
            'Quantity': 2,
            '__effective_user': 'PrimaryForeman',
            '__assignment_method': 'FOREMAN_COLUMN',
            '__is_helper_row': False,
            '__helper_foreman': '',
            '__helper_dept': '',
            '__helper_job': '',
            '__is_vac_crew': False,
            '__source_sheet_id': source_sheet_id or self._SUB_SHEET_ID,
        }

    # ─── Bug A rescue ────────────────────────────────────────────

    def test_bug_a_rescue_returns_expected_price_for_install(self):
        """Bug A rescue: reduced_install × qty for Inst work-type."""
        rescued = generate_weekly_pdfs._subcontractor_rescue_price({
            'CU': 'ANC-M',
            'Work Type': 'Inst',
            'Quantity': 2,
        })
        self.assertEqual(rescued, 100.0)

    def test_bug_a_rescue_decorated_quantity_uses_rate_times_qty(self):
        """2026-08-05 BKT-IP8-F incident: a decorated Quantity ('2 EA')
        must rescue at reduced_install × 2, not fail admission. The
        rescue runs BEFORE row acceptance, so a strict-parse regression
        here drops the row before ``_resolve_row_price`` is reached."""
        rescued = generate_weekly_pdfs._subcontractor_rescue_price({
            'CU': 'ANC-M',
            'Work Type': 'Inst',
            'Quantity': '2 EA',
        })
        self.assertEqual(rescued, 100.0)

    def test_bug_a_rescue_returns_zero_for_unknown_cu(self):
        """Bug A rescue: unknown CU falls through to 0.0 (caller drops row)."""
        rescued = generate_weekly_pdfs._subcontractor_rescue_price({
            'CU': 'UNKNOWN-XYZ',
            'Work Type': 'Inst',
            'Quantity': 2,
        })
        self.assertEqual(rescued, 0.0)

    def test_bug_a_rescue_returns_zero_for_unknown_work_type(self):
        rescued = generate_weekly_pdfs._subcontractor_rescue_price({
            'CU': 'ANC-M',
            'Work Type': 'BogusOp',
            'Quantity': 2,
        })
        self.assertEqual(rescued, 0.0)

    # ─── Bug B1 partitioning ─────────────────────────────────────

    def test_subcontractor_non_helper_row_only_emits_variant_keys(self):
        """Bug B1 partitioning core assertion."""
        row = self._make_synth_non_helper_row(wr='WR_B1')
        groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        # NO legacy primary key
        legacy_primary = '041926_WR_B1'
        self.assertNotIn(
            legacy_primary, keys,
            f"Bug B1 partitioning: subcontractor non-helper row must "
            f"NOT emit legacy primary key; got: {keys}",
        )
        # Variant keys must be present
        self.assertTrue(
            any('REDUCEDSUB' in k and 'HELPER' not in k for k in keys),
            f"Bug B1 must still emit _REDUCEDSUB; got: {keys}",
        )

    def test_non_subcontractor_helper_row_unchanged(self):
        """D-15 scope: non-subcontractor helper rows preserve Phase 1 behavior."""
        row = self._make_synth_helper_row(source_sheet_id=self._NON_SUB_SHEET_ID)
        # Patch lookup_attribution to assert it's never called for non-sub rows
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value={'helper': 'FrozenHelper'},
        ) as mock_lookup:
            groups = generate_weekly_pdfs.group_source_rows([row])
            mock_lookup.assert_not_called()
        keys = list(groups.keys())
        self.assertTrue(
            any('HELPER_ReplacementForeman' in k and 'REDUCEDSUB' not in k for k in keys),
            f"Non-subcontractor helper row should emit legacy "
            f"_HELPER_<name>; got: {keys}",
        )

    # ─── SUB-09 helper-dimension partition (Plan 01.1-06) ────────

    def test_subcontractor_helper_row_does_not_emit_legacy_helper_key(self):
        """Plan 01.1-06 SUB-09 e2e: subcontractor helper row emits NO
        legacy _HELPER_ key (post-cutoff snapshot — both shadow variants).

        Drives group_source_rows with a synthetic sub helper row whose
        snapshot date is 2026-04-19 (post-AEP-cutoff 2026-04-12), so
        BOTH _REDUCEDSUB_HELPER_ and _AEPBILLABLE_HELPER_ are expected.
        The legacy bare-helper key must NOT be present in any form.
        """
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value=None,  # no_history → falls back to current helper
        ):
            row = self._make_synth_helper_row(
                helper_foreman='Drew_Placeholder',
                snapshot='2026-04-19',
            )
            groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        # No legacy _HELPER_ key in any form
        self.assertNotIn(
            '041926_19236776_HELPER_Drew_Placeholder',
            keys,
            f"SUB-09: legacy bare-helper key must NOT be emitted for "
            f"subcontractor helper row; got: {keys}",
        )
        # Confirm no key has variant == 'helper' at all
        for k in keys:
            variant = groups[k][0].get('__variant', '')
            self.assertNotEqual(
                variant, 'helper',
                f"SUB-09: no group should have variant='helper' for a "
                f"subcontractor row; offending key={k!r}, got: {keys}",
            )
        # Shadow variants must be present
        self.assertTrue(
            any('REDUCEDSUB_HELPER_Drew_Placeholder' in k for k in keys),
            f"SUB-09: _REDUCEDSUB_HELPER_ key must be present; got: {keys}",
        )
        self.assertTrue(
            any('AEPBILLABLE_HELPER_Drew_Placeholder' in k for k in keys),
            f"SUB-09: _AEPBILLABLE_HELPER_ key must be present "
            f"(post-cutoff snapshot); got: {keys}",
        )

    def test_subcontractor_helper_row_pre_cutoff_emits_only_reducedsub_helper(self):
        """Plan 01.1-06 SUB-09 e2e: pre-cutoff snapshot — only _REDUCEDSUB_HELPER_
        (mirrors UAT case WR_16087226 wk 041226 which was pre-AEP-cutoff).

        Snapshot 2026-04-11 is BEFORE the AEP cutoff 2026-04-12, so
        _AEPBILLABLE_HELPER_ must NOT be emitted. Legacy _HELPER_ must
        NOT be emitted (SUB-09 partition guard fires regardless of cutoff).
        """
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value=None,  # no_history → falls back to current helper
        ):
            row = self._make_synth_helper_row(
                helper_foreman='Drew_Placeholder',
                snapshot='2026-04-11',  # pre-cutoff
            )
            groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        # _REDUCEDSUB_HELPER_ must be present (unconditional per SUB-02)
        self.assertTrue(
            any('REDUCEDSUB_HELPER_Drew_Placeholder' in k for k in keys),
            f"SUB-09 pre-cutoff: _REDUCEDSUB_HELPER_ must be present; got: {keys}",
        )
        # _AEPBILLABLE_HELPER_ must NOT be present (pre-cutoff snapshot)
        self.assertFalse(
            any('AEPBILLABLE_HELPER_Drew_Placeholder' in k for k in keys),
            f"SUB-09 pre-cutoff: _AEPBILLABLE_HELPER_ must NOT be present; got: {keys}",
        )
        # Legacy _HELPER_ must NOT be present regardless of cutoff
        self.assertFalse(
            any('HELPER_Drew_Placeholder' in k and 'REDUCEDSUB' not in k for k in keys),
            f"SUB-09 pre-cutoff: legacy _HELPER_ must NOT be emitted; got: {keys}",
        )

    # ─── Helper-completed primary exclusion (2026-05-21 hotfix) ──

    def test_subcontractor_helper_row_excluded_from_primary_variant_files(self):
        """Production bug (2026-05-21): a helper-completed subcontractor
        row (both ``Units Completed?`` and ``Helping Foreman Completed
        Unit?`` checked, with a valid ``Foreman Helping?`` + helper dept)
        must NOT be credited to the PRIMARY ``_ReducedSub`` /
        ``_AEPBillable`` files. The helper completed the line item, so it
        belongs SOLELY to the helper-shadow files — otherwise the primary
        foreman is wrongly credited and the line item is double-counted
        (it would appear in both the primary and the helper file). Mirrors
        the legacy main-file ``valid_helper_row`` exclusion.

        Snapshot 2026-04-19 is post-AEP-cutoff so BOTH primary variants
        would be at risk; the assertion proves neither bare primary key
        is emitted while both helper-shadow keys remain.
        """
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value=None,  # no_history → falls back to current helper
        ):
            row = self._make_synth_helper_row(
                helper_foreman='Drew_Placeholder',
                snapshot='2026-04-19',  # post-cutoff
            )
            groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        # NO primary (non-helper) _REDUCEDSUB key for a helper-completed row.
        self.assertFalse(
            any('REDUCEDSUB' in k and 'HELPER' not in k for k in keys),
            f"helper-completed row must NOT emit a primary _REDUCEDSUB "
            f"key (belongs solely to the helper); got: {keys}",
        )
        # NO primary (non-helper) _AEPBILLABLE key for a helper-completed row.
        self.assertFalse(
            any('AEPBILLABLE' in k and 'HELPER' not in k for k in keys),
            f"helper-completed row must NOT emit a primary _AEPBILLABLE "
            f"key (belongs solely to the helper); got: {keys}",
        )
        # Confirm no group carries variant 'reduced_sub' / 'aep_billable'
        # (the primary variants) for this helper row — only the *_helper
        # shadow variants.
        for k in keys:
            variant = groups[k][0].get('__variant', '')
            self.assertNotIn(
                variant, ('reduced_sub', 'aep_billable'),
                f"helper-completed row must not produce a primary variant "
                f"group; offending key={k!r} variant={variant!r}; got: {keys}",
            )
        # The helper-shadow files MUST still be present (helper credit intact).
        self.assertTrue(
            any('REDUCEDSUB_HELPER_Drew_Placeholder' in k for k in keys),
            f"helper-shadow _REDUCEDSUB_HELPER_ must still be present; got: {keys}",
        )
        self.assertTrue(
            any('AEPBILLABLE_HELPER_Drew_Placeholder' in k for k in keys),
            f"helper-shadow _AEPBILLABLE_HELPER_ must still be present; got: {keys}",
        )

    # ─── Bug C attribution ───────────────────────────────────────

    def test_bug_c_attribution_partitions_row_to_frozen_helper(self):
        """Bug C core assertion: frozen helper wins over current helper
        in SHADOW-variant emission (D-15 scope).

        **IN-PLACE REWRITE** (Plan 01.1-06 UAT gap closure, per
        [2026-05-20 00:26] rule 2): the original D-15 assertion
        asserted that the legacy ``_HELPER_ReplacementForeman`` key
        IS still emitted for a subcontractor helper row — encoding the
        buggy additive behavior that Plan 01.1-06 Task 1 closes.

        Under the SUB-09 helper-dimension fix, the legacy helper-key
        append is now gated on ``not is_subcontractor_row``, so a
        subcontractor helper row emits ONLY shadow-variant keys
        (``_REDUCEDSUB_HELPER_<name>`` and ``_AEPBILLABLE_HELPER_<name>``).
        The D-15 assertion is therefore INVERTED: NO legacy
        ``_HELPER_*`` key must be emitted for any subcontractor
        helper row.

        The shadow-variant (Bug C) assertions are PRESERVED unchanged:
        the frozen helper from attribution_snapshot still wins over the
        current Smartsheet value in the shadow files.
        """
        # Phase 2 Plan 02: sub-helper path now calls resolve_claimer with
        # prefetched_map (O(1) map read, D-03). Mock resolve_claimer to
        # return the frozen helper for the 'helper' variant.
        with mock.patch(
            # Store reachable (status reaches resolve_claimer, not the
            # unavailable/fetch_failure short-circuit).
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'no_row'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome('use', 'OriginalForeman', 'frozen', 'success'),
        ):
            row = self._make_synth_helper_row(helper_foreman='ReplacementForeman')
            groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        # Bug C: shadow-variant emission uses the FROZEN helper
        self.assertTrue(
            any('REDUCEDSUB_HELPER_OriginalForeman' in k for k in keys),
            f"Bug C: row's shadow file should use frozen helper; got: {keys}",
        )
        self.assertTrue(
            any('AEPBILLABLE_HELPER_OriginalForeman' in k for k in keys),
            f"Bug C: post-cutoff AEP shadow file should use frozen "
            f"helper; got: {keys}",
        )
        # Bug C: shadow-variant emission MUST NOT use the current helper
        self.assertFalse(
            any('REDUCEDSUB_HELPER_ReplacementForeman' in k for k in keys),
            f"Bug C: shadow file should NOT use current helper; got: {keys}",
        )
        self.assertFalse(
            any('AEPBILLABLE_HELPER_ReplacementForeman' in k for k in keys),
            f"Bug C: shadow file should NOT use current helper; got: {keys}",
        )
        # Plan 01.1-06 SUB-09 helper-dimension fix (D-15 assertion INVERTED):
        # subcontractor helper rows must NOT emit ANY legacy _HELPER_ key.
        # The partition guard (not is_subcontractor_row) means both the
        # current-helper and frozen-helper forms are absent.
        self.assertNotIn(
            '041926_19236776_HELPER_ReplacementForeman',
            keys,
            f"SUB-09 fix: legacy _HELPER_ key must NOT be emitted for "
            f"subcontractor helper row (current helper form); got: {keys}",
        )
        self.assertNotIn(
            '041926_19236776_HELPER_OriginalForeman',
            keys,
            f"SUB-09 fix: legacy _HELPER_ key must NOT be emitted for "
            f"subcontractor helper row (frozen helper form); got: {keys}",
        )

    def test_bug_c_no_history_falls_back_to_current_helper_with_warning(self):
        """D-12 no_history fallback + WARNING discipline.

        Uses the REAL resolve_claimer contract: a genuine no-history row
        returns ResolveOutcome('use', current_value, 'current', 'no_history')
        (writer.py:1048-1049, 1060) — action is 'use', NOT 'no_history'. The
        previous mock returned action='no_history', a value the resolver
        never produces; it exercised an unreachable else branch and hid the
        silent-fallback bug (the 'use' path reset the reason to None, so the
        per-WR WARNING never fired in production).
        """
        with mock.patch(
            # Store reachable + genuinely no frozen row yet (status 'no_row'),
            # so resolve_claimer's no_history is a REAL brand-new claim — not
            # the unavailable short-circuit.
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'no_row'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome(
                'use', 'ReplacementForeman', 'current', 'no_history'),
        ), self.assertLogs(level='WARNING') as log_cm:
            row = self._make_synth_helper_row()
            groups = generate_weekly_pdfs.group_source_rows([row])
        warning_bodies = '\n'.join(log_cm.output)
        self.assertIn(
            'Subcontractor helper claim attribution fallback',
            warning_bodies,
        )
        self.assertIn('reason=no_history', warning_bodies)
        # F1 follow-up: no_history is the benign brand-new-claim case (the
        # lookup SUCCEEDED, just no frozen row yet — this run freezes it), so
        # the remediation must NOT point operators at a Supabase PGRST outage
        # that never happened.
        self.assertIn('No frozen attribution', warning_bodies)
        self.assertNotIn('PGRST', warning_bodies)
        # Row falls back to current helper
        keys = list(groups.keys())
        self.assertTrue(
            any('ReplacementForeman' in k for k in keys),
            f"D-12 fallback: row should be in current helper's file; got: {keys}",
        )

    def test_bug_c_fetch_failure_falls_back_with_correct_reason(self):
        """Distinguish no_history vs fetch_failure per D-12.

        Phase 2 Plan 02 update: the sub-helper path now calls
        resolve_claimer with prefetched_map (O(1) map read, D-03).
        A 'hold' outcome (action='hold') signals fetch_failure and
        triggers the D-12 fallback WARNING with reason=fetch_failure.
        """
        with mock.patch(
            # Store reachable (status 'no_row' reaches resolve_claimer); the
            # 'hold' outcome is what signals fetch_failure here — distinct from
            # the store-level 'unavailable' short-circuit.
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'no_row'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome('hold', None, None, 'fetch_failure'),
        ), self.assertLogs(level='WARNING') as log_cm:
            row = self._make_synth_helper_row()
            generate_weekly_pdfs.group_source_rows([row])
        warning_bodies = '\n'.join(log_cm.output)
        self.assertIn('reason=fetch_failure', warning_bodies)
        # fetch_failure IS a real PostgREST outage — keep the Supabase Logs
        # investigation guidance (contrast with the benign no_history case).
        self.assertIn('PGRST', warning_bodies)

    def test_bug_c_warning_dedupe_per_wr_helper(self):
        """Per-WR WARNING fires ONCE per (wr, week, helper) tuple.

        Uses the REAL resolve_claimer contract (action='use',
        reason='no_history') so the dedupe is verified on the path
        production actually takes, not the unreachable else branch.
        """
        with mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome(
                'use', 'ReplacementForeman', 'current', 'no_history'),
        ), self.assertLogs(level='WARNING') as log_cm:
            rows = [
                self._make_synth_helper_row(row_id=i)
                for i in range(100, 105)
            ]
            generate_weekly_pdfs.group_source_rows(rows)
        fallback_warnings = [
            line for line in log_cm.output
            if 'Subcontractor helper claim attribution fallback' in line
        ]
        self.assertEqual(
            len(fallback_warnings), 1,
            f"Per-WR dedupe broken: expected 1 WARNING, got "
            f"{len(fallback_warnings)}: {fallback_warnings}",
        )

    def test_bug_c_kill_switch_off_uses_current_helper(self):
        """Bug C kill switch — flips to D-12 unconditional fallback."""
        generate_weekly_pdfs.SUBCONTRACTOR_HELPER_CLAIM_ATTRIBUTION_ENABLED = False
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value={'helper': 'OriginalForeman'},
        ) as mock_lookup:
            row = self._make_synth_helper_row()
            groups = generate_weekly_pdfs.group_source_rows([row])
            mock_lookup.assert_not_called()
        keys = list(groups.keys())
        self.assertTrue(
            any('REDUCEDSUB_HELPER_ReplacementForeman' in k for k in keys),
            f"Bug C kill-switch-off: row should go to current helper; "
            f"got: {keys}",
        )


class TestRpcMissingGracefulDegradation(unittest.TestCase):
    """Phase 2 Plan 05 (CR-01): a MISSING bulk RPC (rpc_missing) degrades to
    the per-row path instead of HOLDing every B/C/sub-helper row.

    Mirrors TestEndToEndPipeline's module-state setup (subcontractor sheet,
    variant + attribution flags enabled, seeded rate) and reuses its
    ``_make_synth_*`` row builders via delegation. Each test mocks
    ``prefetch_attribution`` directly to control _attr_status, and
    ``resolve_claimer`` to make the per-row outcome deterministic.
    """

    _SUB_SHEET_ID = TestEndToEndPipeline._SUB_SHEET_ID

    def setUp(self):
        _reset_all()
        _ensure_smartsheet_mocked()
        self._orig_enabled = generate_weekly_pdfs.SUBCONTRACTOR_RATE_VARIANTS_ENABLED
        self._orig_bug_a = generate_weekly_pdfs.SUBCONTRACTOR_RATE_RECALC_PREACCEPTANCE_ENABLED
        self._orig_bug_c = generate_weekly_pdfs.SUBCONTRACTOR_HELPER_CLAIM_ATTRIBUTION_ENABLED
        self._orig_sub_ids = set(generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS)
        self._orig_rates = dict(generate_weekly_pdfs._SUBCONTRACTOR_RATES)
        self._orig_fallback = generate_weekly_pdfs.ATTRIBUTION_BULK_PREFETCH_FALLBACK
        generate_weekly_pdfs.SUBCONTRACTOR_RATE_VARIANTS_ENABLED = True
        generate_weekly_pdfs.SUBCONTRACTOR_RATE_RECALC_PREACCEPTANCE_ENABLED = True
        generate_weekly_pdfs.SUBCONTRACTOR_HELPER_CLAIM_ATTRIBUTION_ENABLED = True
        generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS.clear()
        generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS.add(self._SUB_SHEET_ID)
        generate_weekly_pdfs._SUBCONTRACTOR_RATES['ANC-M'] = {
            'new_install_price': 75.0,
            'reduced_install_price': 50.0,
            'new_remove_price': 60.0,
            'reduced_remove_price': 45.0,
            'new_transfer_price': 80.0,
            'reduced_transfer_price': 55.0,
        }

    def tearDown(self):
        generate_weekly_pdfs.SUBCONTRACTOR_RATE_VARIANTS_ENABLED = self._orig_enabled
        generate_weekly_pdfs.SUBCONTRACTOR_RATE_RECALC_PREACCEPTANCE_ENABLED = self._orig_bug_a
        generate_weekly_pdfs.SUBCONTRACTOR_HELPER_CLAIM_ATTRIBUTION_ENABLED = self._orig_bug_c
        generate_weekly_pdfs.ATTRIBUTION_BULK_PREFETCH_FALLBACK = self._orig_fallback
        generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS.clear()
        generate_weekly_pdfs._FOLDER_DISCOVERED_SUB_IDS.update(self._orig_sub_ids)
        generate_weekly_pdfs._SUBCONTRACTOR_RATES.clear()
        generate_weekly_pdfs._SUBCONTRACTOR_RATES.update(self._orig_rates)
        _reset_all()

    # Reuse the parent's synthetic-row builders without inheriting its tests.
    _make_synth_helper_row = TestEndToEndPipeline._make_synth_helper_row
    _make_synth_non_helper_row = TestEndToEndPipeline._make_synth_non_helper_row

    def test_rpc_missing_with_fallback_generates_sub_helper(self):
        """rpc_missing + fallback ON: sub-helper resolves per-row and the
        row GENERATES (no HOLD, no fetch_failure WARNING)."""
        generate_weekly_pdfs.ATTRIBUTION_BULK_PREFETCH_FALLBACK = True
        with mock.patch(
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'rpc_missing'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome('use', 'FrozenHelper', 'frozen', 'success'),
        ) as _rc:
            row = self._make_synth_helper_row()
            groups = generate_weekly_pdfs.group_source_rows([row])
        # Per-row resolver WAS consulted (degrade path), not bypassed.
        self.assertTrue(_rc.called)
        keys = list(groups.keys())
        self.assertTrue(
            any('REDUCEDSUB_HELPER_FrozenHelper' in k for k in keys),
            f"rpc_missing fallback: row should generate under the frozen "
            f"helper; got: {keys}",
        )

    def test_rpc_missing_with_fallback_b_generates_user_variant(self):
        """rpc_missing + fallback ON: subcontractor non-helper PRIMARY row
        generates a _REDUCEDSUB_USER_ group (no HOLD)."""
        generate_weekly_pdfs.ATTRIBUTION_BULK_PREFETCH_FALLBACK = True
        with mock.patch(
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'rpc_missing'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome('use', 'FrozenPrimary', 'frozen', 'success'),
        ):
            row = self._make_synth_non_helper_row()
            row['__row_id'] = 71001
            groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        self.assertTrue(
            any('REDUCEDSUB_USER_FrozenPrimary' in k for k in keys),
            f"rpc_missing fallback: B should emit a per-claimer primary "
            f"group; got: {keys}",
        )

    def test_fetch_failure_still_holds_b_no_user_variant(self):
        """fetch_failure: B HOLDs (D-04) — no _REDUCEDSUB_USER_ group emitted,
        and resolve_claimer is NOT consulted (direct HOLD, zero RPC)."""
        with mock.patch(
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'fetch_failure'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome('use', 'ShouldNotAppear', 'frozen', 'success'),
        ) as _rc:
            row = self._make_synth_non_helper_row()
            row['__row_id'] = 72001
            groups = generate_weekly_pdfs.group_source_rows([row])
        _rc.assert_not_called()
        keys = list(groups.keys())
        self.assertFalse(
            any('REDUCEDSUB_USER_' in k for k in keys),
            f"fetch_failure must HOLD the B row (no _USER_ group); got: {keys}",
        )

    def test_rpc_missing_fallback_off_holds_b(self):
        """rpc_missing + fallback OFF: operator opted out -> B HOLDs (no
        _REDUCEDSUB_USER_ group)."""
        generate_weekly_pdfs.ATTRIBUTION_BULK_PREFETCH_FALLBACK = False
        with mock.patch(
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'rpc_missing'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome('use', 'ShouldNotAppear', 'frozen', 'success'),
        ) as _rc:
            row = self._make_synth_non_helper_row()
            row['__row_id'] = 73001
            groups = generate_weekly_pdfs.group_source_rows([row])
        _rc.assert_not_called()
        keys = list(groups.keys())
        self.assertFalse(
            any('REDUCEDSUB_USER_' in k for k in keys),
            f"fallback-off rpc_missing must HOLD the B row; got: {keys}",
        )

    def test_wr05_fetch_failure_sub_helper_emits_warning(self):
        """WR-05: on fetch_failure the sub-helper block sets
        _attribution_reason='fetch_failure' so the per-WR WARNING fires,
        WITHOUT consulting the per-row resolver (direct status thread)."""
        with mock.patch(
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'fetch_failure'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome('use', 'X', 'frozen', 'success'),
        ) as _rc, self.assertLogs(level='WARNING') as log_cm:
            row = self._make_synth_helper_row()
            generate_weekly_pdfs.group_source_rows([row])
        warning_bodies = '\n'.join(log_cm.output)
        self.assertIn('reason=fetch_failure', warning_bodies)
        self.assertIn(
            'Subcontractor helper claim attribution fallback', warning_bodies
        )
        # WR-05 threads the status directly; the per-row resolver is NOT
        # re-invoked for the strict-HOLD case.
        _rc.assert_not_called()

    def test_wr05_rpc_missing_fallback_off_sub_helper_emits_warning(self):
        """WR-05: rpc_missing + fallback OFF also surfaces the WARNING."""
        generate_weekly_pdfs.ATTRIBUTION_BULK_PREFETCH_FALLBACK = False
        with mock.patch(
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'rpc_missing'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome('use', 'X', 'frozen', 'success'),
        ), self.assertLogs(level='WARNING') as log_cm:
            row = self._make_synth_helper_row()
            generate_weekly_pdfs.group_source_rows([row])
        warning_bodies = '\n'.join(log_cm.output)
        self.assertIn('reason=fetch_failure', warning_bodies)

    def test_wr05_unavailable_sub_helper_emits_distinct_warning(self):
        """Codex P2 (PR #281): 'unavailable' (no Supabase client) must NOT be
        collapsed to 'no_history'. It threads the status directly (resolver not
        consulted) and gets a config-oriented remediation — never the benign
        no_history "this run freezes it; no action needed" text, and never the
        fetch_failure PGRST guidance."""
        with mock.patch(
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'unavailable'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome('use', 'X', 'frozen', 'success'),
        ) as _rc, self.assertLogs(level='WARNING') as log_cm:
            row = self._make_synth_helper_row()
            generate_weekly_pdfs.group_source_rows([row])
        warning_bodies = '\n'.join(log_cm.output)
        self.assertIn('reason=unavailable', warning_bodies)
        self.assertIn(
            'Subcontractor helper claim attribution fallback', warning_bodies
        )
        # Config-oriented remediation, NOT the benign no_history message.
        self.assertIn('unavailable', warning_bodies)
        self.assertIn('SUPABASE', warning_bodies)
        self.assertNotIn('this run freezes it', warning_bodies)
        self.assertNotIn('reason=no_history', warning_bodies)
        self.assertNotIn('PGRST', warning_bodies)
        # The sub-helper block threads the status directly — resolve_claimer is
        # NOT consulted for the 'helper' role. (The reduced_sub/B primary path
        # legitimately calls it for its own role with use-current fallback; that
        # availability-first behavior is unrelated and intentionally unchanged.)
        helper_role_calls = [
            c for c in _rc.call_args_list if c.args and c.args[0] == 'helper'
        ]
        self.assertEqual(
            helper_role_calls, [],
            "sub-helper block must thread 'unavailable' directly, not consult "
            "resolve_claimer for the 'helper' role",
        )


class TestBugB2WhitelistE2E(unittest.TestCase):
    """D-21(c): PPP cleanup whitelist defense-in-depth."""

    def setUp(self):
        _ensure_smartsheet_mocked()

    def _make_attachment(self, name, att_id):
        att = mock.MagicMock()
        att.name = name
        att.id = att_id
        return att

    def _build_client_with_attachments(self, attachments):
        client = mock.MagicMock()
        sheet = mock.MagicMock()
        row = mock.MagicMock()
        row.id = 1
        client.Attachments.list_row_attachments.return_value.data = attachments
        sheet.rows = [row]
        client.Sheets.get_sheet.return_value = sheet
        return client, sheet

    def test_ppp_off_contract_primary_attachment_deleted(self):
        """Off-contract variant on PPP is unconditionally deleted."""
        att_primary = self._make_attachment(
            'WR_19236776_WeekEnding_041926_120000_abc123.xlsx', 10
        )
        att_reduced = self._make_attachment(
            'WR_19236776_WeekEnding_041926_120000_ReducedSub_def456.xlsx', 20
        )
        client, sheet = self._build_client_with_attachments(
            [att_primary, att_reduced]
        )
        generate_weekly_pdfs.cleanup_untracked_sheet_attachments(
            client,
            target_sheet_id=8162920222379908,
            valid_wr_weeks=set(),
            test_mode=False,
            target_sheet=sheet,
            variant_whitelist={'reduced_sub', 'reduced_sub_helper'},
        )
        deletes = [
            call.args for call in client.Attachments.delete_attachment.call_args_list
        ]
        self.assertIn(
            (8162920222379908, 10), deletes,
            f"primary attachment must be deleted (off-contract for PPP); "
            f"got deletes={deletes}"
        )
        # reduced_sub attachment should NOT be deleted as off-contract.
        # It IS in identity_groups; with empty valid_wr_weeks and the
        # default KEEP_HISTORICAL_WEEKS=False, the legacy path may
        # still delete it via the variant-keep-newest cleanup OR may
        # leave it (depends on whether it's a "newest" or "older"
        # variant in the single-element identity group; single-element
        # groups always have only one entry so atts_sorted[1:] is empty
        # and no delete fires). So the reduced_sub MUST NOT be in the
        # off-contract delete batch even though it may be deleted by
        # other paths — the assertion focuses on the off-contract
        # invariant.
        # In fact for this single-attachment identity group there's
        # nothing to delete, so verify the only delete was the
        # off-contract primary:
        self.assertEqual(
            len(deletes), 1,
            f"Only the off-contract primary attachment should be "
            f"deleted; got: {deletes}"
        )

    def test_target_cleanup_with_none_whitelist_preserves_legacy(self):
        """TARGET cleanup with variant_whitelist=None preserves Phase 1 behavior."""
        att_primary = self._make_attachment(
            'WR_19236776_WeekEnding_041926_120000_abc123.xlsx', 10
        )
        client, sheet = self._build_client_with_attachments([att_primary])
        generate_weekly_pdfs.cleanup_untracked_sheet_attachments(
            client,
            target_sheet_id=5723337641643908,
            valid_wr_weeks={('19236776', '041926', 'primary', '')},
            test_mode=False,
            target_sheet=sheet,
            variant_whitelist=None,
        )
        deletes = [
            call.args for call in client.Attachments.delete_attachment.call_args_list
        ]
        # With whitelist=None, the off-contract loop is skipped entirely.
        # The single-attachment identity group has no older variants to
        # delete (atts_sorted[1:] is empty), so no deletes fire.
        self.assertEqual(
            deletes, [],
            f"TARGET with whitelist=None should not unconditionally "
            f"delete; got: {deletes}"
        )

    def test_ppp_with_only_whitelisted_attachments_no_delete(self):
        """PPP with only whitelisted variants performs zero off-contract deletes."""
        att_reduced = self._make_attachment(
            'WR_19236776_WeekEnding_041926_120000_ReducedSub_def456.xlsx', 20
        )
        client, sheet = self._build_client_with_attachments([att_reduced])
        generate_weekly_pdfs.cleanup_untracked_sheet_attachments(
            client,
            target_sheet_id=8162920222379908,
            valid_wr_weeks=set(),
            test_mode=False,
            target_sheet=sheet,
            variant_whitelist={'reduced_sub', 'reduced_sub_helper'},
        )
        deletes = [
            call.args for call in client.Attachments.delete_attachment.call_args_list
        ]
        self.assertEqual(
            deletes, [],
            f"PPP with only whitelisted variant attachment must not "
            f"trigger off-contract delete; got: {deletes}"
        )


class TestLegacyHelperTargetCleanupE2E(unittest.TestCase):
    """Plan 01.1-06 SUB-09: TARGET cleanup removes pre-existing legacy
    _Helper_ and bare-primary attachments for subcontractor WRs via the
    new sub_wr_scope / sub_offcontract_variants kwargs.
    """

    def setUp(self):
        _ensure_smartsheet_mocked()

    def _make_attachment(self, name, att_id):
        att = mock.MagicMock()
        att.name = name
        att.id = att_id
        return att

    def _build_client_with_attachments(self, attachments):
        client = mock.MagicMock()
        sheet = mock.MagicMock()
        row = mock.MagicMock()
        row.id = 1
        client.Attachments.list_row_attachments.return_value.data = attachments
        sheet.rows = [row]
        client.Sheets.get_sheet.return_value = sheet
        return client, sheet

    def test_target_cleanup_removes_legacy_helper_for_subcontractor_wr(self):
        """TARGET cleanup deletes _Helper_ attachment for sub WR via sub_wr_scope."""
        att_helper = self._make_attachment(
            'WR_16087226_WeekEnding_041226_220404_Helper_Drew_Placeholder_abc123.xlsx',
            10,
        )
        att_shadow = self._make_attachment(
            'WR_16087226_WeekEnding_041226_220404_ReducedSub_Helper_Drew_Placeholder_def456.xlsx',
            20,
        )
        client, sheet = self._build_client_with_attachments(
            [att_helper, att_shadow]
        )
        generate_weekly_pdfs.cleanup_untracked_sheet_attachments(
            client,
            target_sheet_id=5723337641643908,
            valid_wr_weeks=set(),
            test_mode=False,
            target_sheet=sheet,
            sub_wr_scope={'16087226'},
            sub_offcontract_variants={'helper', 'primary'},
        )
        deletes = [call.args for call in client.Attachments.delete_attachment.call_args_list]
        self.assertIn(
            (5723337641643908, 10), deletes,
            f"_Helper_ attachment must be deleted as off-contract for sub WR; "
            f"got deletes={deletes}",
        )
        # _ReducedSub_Helper_ is on-contract for sub WR — must NOT be deleted
        # by the sub_wr_scope path (it lands in identity_groups and may be
        # pruned only as an "older variant" by the keep-newest logic, but with
        # a single-attachment identity group atts_sorted[1:] is empty, so no
        # delete fires on that path either).
        self.assertNotIn(
            (5723337641643908, 20), deletes,
            f"_ReducedSub_Helper_ shadow must NOT be deleted; got: {deletes}",
        )

    def test_target_cleanup_preserves_legacy_helper_for_non_sub_wr(self):
        """TARGET cleanup does NOT delete _Helper_ attachment for non-sub WR.

        WR 10277491 is NOT in sub_wr_scope {'16087226'} so the new gate
        is a no-op and byte-identical legacy TARGET behaviour is preserved.
        """
        att_helper = self._make_attachment(
            'WR_10277491_WeekEnding_041226_220404_Helper_SomeForeman_abc123.xlsx',
            30,
        )
        client, sheet = self._build_client_with_attachments([att_helper])
        generate_weekly_pdfs.cleanup_untracked_sheet_attachments(
            client,
            target_sheet_id=5723337641643908,
            valid_wr_weeks={('10277491', '041226', 'helper', 'SomeForeman')},
            test_mode=False,
            target_sheet=sheet,
            sub_wr_scope={'16087226'},  # different WR — 10277491 not in scope
            sub_offcontract_variants={'helper', 'primary'},
        )
        deletes = [call.args for call in client.Attachments.delete_attachment.call_args_list]
        self.assertNotIn(
            (5723337641643908, 30), deletes,
            f"Non-sub WR _Helper_ must NOT be deleted; got: {deletes}",
        )

    def test_target_cleanup_exempts_live_helper_for_overlapping_sub_wr(self):
        """WR-01 (review follow-up): cross-sheet WR overlap exemption.

        A single WR# can be in ``sub_wr_scope`` (it has subcontractor
        ``_ReducedSub`` rows on a sub sheet) AND simultaneously have a
        LEGITIMATE live ``_Helper_<name>.xlsx`` produced this run from a
        NON-subcontractor sheet (its identity is in ``valid_wr_weeks``).
        Because ``sub_wr_scope`` keys on WR# alone but
        ``is_subcontractor_row`` is decided per-row by source sheet, the
        SUB-09 off-contract gate would otherwise delete that live file on
        every run → regenerate → re-upload → delete (a churn loop with a
        data-absent window). The ``ident not in valid_wr_weeks`` guard
        must EXEMPT the live identity while STILL deleting a stale orphan
        legacy ``_Helper_`` (different identity, NOT in valid_wr_weeks) for
        the same in-scope WR.
        """
        # Live non-sub helper for the in-scope WR — identity IS in
        # valid_wr_weeks (parses to ('16087226','041226','helper','Live_Foreman')).
        att_live = self._make_attachment(
            'WR_16087226_WeekEnding_041226_220404_Helper_Live_Foreman_abc123.xlsx',
            40,
        )
        # Stale orphan legacy helper for the SAME in-scope WR — identity
        # ('16087226','041226','helper','Old_Foreman') is NOT in valid_wr_weeks.
        att_orphan = self._make_attachment(
            'WR_16087226_WeekEnding_041226_220404_Helper_Old_Foreman_def456.xlsx',
            50,
        )
        client, sheet = self._build_client_with_attachments(
            [att_live, att_orphan]
        )
        generate_weekly_pdfs.cleanup_untracked_sheet_attachments(
            client,
            target_sheet_id=5723337641643908,
            valid_wr_weeks={('16087226', '041226', 'helper', 'Live_Foreman')},
            test_mode=False,
            target_sheet=sheet,
            sub_wr_scope={'16087226'},
            sub_offcontract_variants={'helper', 'primary'},
        )
        deletes = [
            call.args for call in client.Attachments.delete_attachment.call_args_list
        ]
        self.assertNotIn(
            (5723337641643908, 40), deletes,
            f"WR-01: a live non-sub _Helper_ whose identity is in "
            f"valid_wr_weeks must be EXEMPT from the sub_wr_scope "
            f"off-contract gate; got deletes={deletes}",
        )
        self.assertIn(
            (5723337641643908, 50), deletes,
            f"WR-01: a stale orphan _Helper_ (NOT in valid_wr_weeks) for "
            f"the same in-scope WR must still be deleted as off-contract; "
            f"got deletes={deletes}",
        )


class TestHashPruneIdempotency(unittest.TestCase):
    """D-21(e) + SUB-12 + Pitfall 4 closure."""

    def setUp(self):
        _ensure_smartsheet_mocked()

    def _make_groups_with_reducedsub(self, wrs):
        """Build a synthetic groups dict containing _REDUCEDSUB-suffixed
        keys for the given WRs — drives ``_run_phase_1_1_hash_prune``'s
        simplified D-18 detection."""
        groups = {}
        for wr in wrs:
            key = f"041926_{wr}_REDUCEDSUB"
            groups[key] = [{
                'Work Request #': wr,
                '__source_sheet_id': 8162920222379908,
                # Production rows always carry __variant (set at emission);
                # the scope builder now gates on it.
                '__variant': 'reduced_sub',
            }]
        return groups

    # ─── Pitfall 4 sentinel round-trip ───────────────────────────
    # Phase 11 Plan 08 (INC-05, D-12) REWRITES this section in place of
    # the retired round-trip tests: load_hash_history / save_hash_history
    # / HASH_HISTORY_MAX_ENTRIES are removed along with
    # generated_docs/hash_history.json itself -- there is no longer a
    # JSON round trip (or int-sentinel retention-sort) to pin.
    # _run_phase_1_1_hash_prune itself stays defined in
    # pipeline/attribution.py (harmless, uncalled) and is still exercised
    # directly against an in-memory fixture dict by the "Prune behavior"
    # tests below.

    def test_hash_history_persistence_helpers_removed(self):
        for name in ('load_hash_history', 'save_hash_history',
                     'HASH_HISTORY_MAX_ENTRIES'):
            self.assertFalse(
                hasattr(generate_weekly_pdfs, name),
                f"{name} should have been removed (Phase 11 Plan 08 / INC-05)",
            )

    # ─── Prune behavior ──────────────────────────────────────────

    def test_first_run_advances_version_and_drops_orphans(self):
        """Version 0 → 1: orphans dropped, sentinel persists, log fires."""
        hash_history = {
            '19236776|041926|primary|': {'hash': 'h1', 'timestamp': '2026-01-01'},
            '13478502|041926|primary|': {'hash': 'h2', 'timestamp': '2026-01-02'},
            '15613821|041926|primary|': {'hash': 'h3', 'timestamp': '2026-01-03'},
            # Unaffected entries — non-subcontractor WR
            '12345|041926|primary|': {'hash': 'h4', 'timestamp': '2026-01-04'},
            # Unaffected entry — non-primary variant for a subcontractor WR
            '19236776|041926|reduced_sub|': {
                'hash': 'h5', 'timestamp': '2026-01-05',
            },
        }
        groups = self._make_groups_with_reducedsub(
            ['19236776', '13478502', '15613821']
        )
        with self.assertLogs(level='INFO') as log_cm:
            generate_weekly_pdfs._run_phase_1_1_hash_prune(hash_history, groups)
        # 3 orphans dropped
        self.assertNotIn('19236776|041926|primary|', hash_history)
        self.assertNotIn('13478502|041926|primary|', hash_history)
        self.assertNotIn('15613821|041926|primary|', hash_history)
        # Unaffected entries preserved
        self.assertIn('12345|041926|primary|', hash_history)
        self.assertIn('19236776|041926|reduced_sub|', hash_history)
        # Sentinel persists
        self.assertEqual(
            hash_history['_phase_prune_version'],
            generate_weekly_pdfs.PHASE_1_1_HASH_PRUNE_VERSION,
        )
        # ONE info log mentioning the dropped count
        prune_logs = [
            line for line in log_cm.output
            if 'Phase 1.1 hash-history prune' in line
        ]
        self.assertEqual(len(prune_logs), 1, f"Expected 1 prune log; got: {prune_logs}")
        self.assertIn('dropped 3', prune_logs[0])

    def test_subsequent_run_at_current_version_is_noop(self):
        """Sentinel already at current version → no-op (no entries dropped)."""
        hash_history = {
            '19236776|041926|primary|': {'hash': 'h1', 'timestamp': '2026-01-01'},
            '_phase_prune_version': generate_weekly_pdfs.PHASE_1_1_HASH_PRUNE_VERSION,
        }
        groups = self._make_groups_with_reducedsub(['19236776'])
        generate_weekly_pdfs._run_phase_1_1_hash_prune(hash_history, groups)
        # Orphan preserved (no-op path)
        self.assertIn('19236776|041926|primary|', hash_history)
        # Sentinel preserved at current value
        self.assertEqual(
            hash_history['_phase_prune_version'],
            generate_weekly_pdfs.PHASE_1_1_HASH_PRUNE_VERSION,
        )

    def test_prune_excludes_non_subcontractor_primary(self):
        """D-18 scope: only WRs whose groups contain _REDUCEDSUB are in-scope."""
        hash_history = {
            '99999|041926|primary|': {'hash': 'h1', 'timestamp': '2026-01-01'},
        }
        # groups dict has NO _REDUCEDSUB key for WR 99999
        groups = self._make_groups_with_reducedsub(['some_other_wr'])
        generate_weekly_pdfs._run_phase_1_1_hash_prune(hash_history, groups)
        # 99999 entry preserved (out of scope)
        self.assertIn('99999|041926|primary|', hash_history)

    def test_prune_excludes_non_primary_subcontractor_variants(self):
        """Version 2 scope: only 'primary' (4-part blank-id) and 'helper' orphans
        are dropped for in-scope sub WRs — shadow variants are preserved.

        **IN-PLACE REWRITE** (Plan 01.1-06, per [2026-05-20 00:26] rule 2):
        under v2 the prune now ALSO drops 4-part 'helper' keys (variant='helper')
        for in-scope sub WRs. The original v1 assertion that
        '19236776|041926|helper|Foreman' IS retained is INVERTED. The live
        shadow variants (reduced_sub, aep_billable, reduced_sub_helper) are
        STILL retained and their assertIn calls are PRESERVED unchanged.
        """
        hash_history = {
            '19236776|041926|reduced_sub|': {
                'hash': 'h1', 'timestamp': '2026-01-01',
            },
            '19236776|041926|aep_billable|': {
                'hash': 'h2', 'timestamp': '2026-01-02',
            },
            '19236776|041926|reduced_sub_helper|Foreman': {
                'hash': 'h3', 'timestamp': '2026-01-03',
            },
            '19236776|041926|helper|Foreman': {
                'hash': 'h4', 'timestamp': '2026-01-04',
            },
        }
        groups = self._make_groups_with_reducedsub(['19236776'])
        generate_weekly_pdfs._run_phase_1_1_hash_prune(hash_history, groups)
        # Shadow variants must still be retained (UNCHANGED from v1 assertions)
        self.assertIn('19236776|041926|reduced_sub|', hash_history)
        self.assertIn('19236776|041926|aep_billable|', hash_history)
        self.assertIn('19236776|041926|reduced_sub_helper|Foreman', hash_history)
        # v2 INVERTS the helper assertion: a 4-part 'helper' key for an in-scope
        # sub WR is NOW DROPPED (it is a legacy orphan from pre-Plan-01.1-06 runs)
        self.assertNotIn(
            '19236776|041926|helper|Foreman',
            hash_history,
            "v2 prune must drop 4-part 'helper' key for in-scope sub WR",
        )

    def test_reset_hash_history_followed_by_prune_is_noop(self):
        """RESET_HASH_HISTORY=true → empty dict → prune writes sentinel + 0 drops."""
        hash_history = {}  # simulates load_hash_history after RESET
        groups = self._make_groups_with_reducedsub(['19236776'])
        with self.assertLogs(level='INFO') as log_cm:
            generate_weekly_pdfs._run_phase_1_1_hash_prune(hash_history, groups)
        # 0 entries dropped (nothing was there), sentinel persists
        self.assertEqual(
            hash_history['_phase_prune_version'],
            generate_weekly_pdfs.PHASE_1_1_HASH_PRUNE_VERSION,
        )
        # ONE info log — "no primary/legacy-helper orphans to drop" path
        # (wording updated in Plan 01.1-06 Task 3 to reflect v2 scope)
        prune_logs = [
            line for line in log_cm.output
            if 'Phase 1.1 hash-history prune' in line
        ]
        self.assertEqual(len(prune_logs), 1)
        self.assertIn('no primary/legacy-helper orphans to drop', prune_logs[0])

    def test_first_run_with_no_orphans_logs_no_orphan_branch(self):
        """Version 0 + groups with _REDUCEDSUB + no matching orphans → log + noop."""
        hash_history = {
            'other_wr|041926|primary|': {'hash': 'h1', 'timestamp': '2026-01-01'},
        }
        groups = self._make_groups_with_reducedsub(['19236776'])
        with self.assertLogs(level='INFO') as log_cm:
            generate_weekly_pdfs._run_phase_1_1_hash_prune(hash_history, groups)
        # Non-subcontractor entry preserved
        self.assertIn('other_wr|041926|primary|', hash_history)
        # Sentinel persists
        self.assertEqual(
            hash_history['_phase_prune_version'],
            generate_weekly_pdfs.PHASE_1_1_HASH_PRUNE_VERSION,
        )
        prune_logs = [
            line for line in log_cm.output
            if 'Phase 1.1 hash-history prune' in line
        ]
        self.assertEqual(len(prune_logs), 1)
        # Wording updated in Plan 01.1-06 Task 3 to reflect v2 scope
        self.assertIn('no primary/legacy-helper orphans to drop', prune_logs[0])


    def test_version_2_drops_subcontractor_legacy_helper_orphans(self):
        """Plan 01.1-06 Task 3: version-2 prune drops 6-part 'helper' orphans
        for in-scope sub WRs, AND preserves the version-1 primary-orphan drop,
        AND preserves non-sub helper entries and live shadow entries.

        The 6-part key shape exercises the Task 3 parse fix: the former
        '!= 4' guard hard-skipped every helper key, so this test WOULD HAVE
        FAILED before Task 3 (the 6-part orphan would survive the prune).
        """
        hash_history = {
            # 6-part subcontractor helper orphan (THE KEY CASE)
            '16087226|041226|helper|Drew_Placeholder|500|JOB-A': {
                'hash': 'h1', 'timestamp': '2026-01-01',
            },
            # 4-part subcontractor primary orphan (version-1 case)
            '16087226|041226|primary|': {
                'hash': 'h2', 'timestamp': '2026-01-02',
            },
            # 6-part NON-sub helper (wr not in scope — must survive)
            '10277491|041226|helper|Other|600|JOB-B': {
                'hash': 'h3', 'timestamp': '2026-01-03',
            },
            # Live shadow variant (must survive)
            '16087226|041226|reduced_sub_helper|Drew_Placeholder': {
                'hash': 'h4', 'timestamp': '2026-01-04',
            },
            '_phase_prune_version': 1,
        }
        groups = self._make_groups_with_reducedsub(['16087226'])
        with self.assertLogs(level='INFO') as log_cm:
            generate_weekly_pdfs._run_phase_1_1_hash_prune(hash_history, groups)
        # 6-part sub helper orphan DROPPED
        self.assertNotIn(
            '16087226|041226|helper|Drew_Placeholder|500|JOB-A',
            hash_history,
            "v2: 6-part sub helper orphan must be dropped",
        )
        # 4-part sub primary orphan DROPPED (version-1 superset)
        self.assertNotIn(
            '16087226|041226|primary|',
            hash_history,
            "v2: 4-part sub primary orphan must also be dropped",
        )
        # 6-part NON-sub helper PRESERVED
        self.assertIn(
            '10277491|041226|helper|Other|600|JOB-B',
            hash_history,
            "v2: non-sub helper entry must be preserved",
        )
        # Live shadow PRESERVED
        self.assertIn(
            '16087226|041226|reduced_sub_helper|Drew_Placeholder',
            hash_history,
            "v2: live shadow variant must be preserved",
        )
        # Sentinel advanced to 2
        self.assertEqual(
            hash_history['_phase_prune_version'],
            generate_weekly_pdfs.PHASE_1_1_HASH_PRUNE_VERSION,
        )
        # Log must mention dropped count
        prune_logs = [
            line for line in log_cm.output
            if 'Phase 1.1 hash-history prune' in line
        ]
        self.assertEqual(len(prune_logs), 1)
        self.assertIn('dropped 2', prune_logs[0])

    def test_version_2_idempotent_when_sentinel_already_2(self):
        """Sentinel already at 2 → no-op (idempotency per [2026-04-25 12:00] rule 1)."""
        hash_history = {
            '16087226|041226|helper|Drew_Placeholder|500|JOB-A': {
                'hash': 'h1', 'timestamp': '2026-01-01',
            },
            '_phase_prune_version': 2,
        }
        groups = self._make_groups_with_reducedsub(['16087226'])
        generate_weekly_pdfs._run_phase_1_1_hash_prune(hash_history, groups)
        # No-op: helper orphan preserved (already migrated)
        self.assertIn(
            '16087226|041226|helper|Drew_Placeholder|500|JOB-A',
            hash_history,
            "Idempotency: no drops when sentinel is already at current version",
        )
        # Sentinel preserved
        self.assertEqual(hash_history['_phase_prune_version'], 2)


class TestProductionCodeSiteInvariants(unittest.TestCase):
    """Source-level grep guards — defeat the 'mirror passes but
    production reverted' failure mode that allowed Phase 1 to ship
    with Bugs A and B1 latent.
    """

    @classmethod
    def setUpClass(cls):
        # Phase 09 W3: the Bug-A rescue gate (is_subcontractor_sheet +
        # SUBCONTRACTOR_RATE_RECALC_PREACCEPTANCE_ENABLED + _subcontractor_
        # rescue_price) lives inside get_all_source_rows, relocated to
        # pipeline/fetch.py. Grep the facade + the relocated module so the
        # production-site invariants follow the code.
        import pipeline.fetch
        import pipeline.grouping  # W4: group_source_rows relocated here
        import pipeline.cleanup  # W5: cleanup_untracked_sheet_attachments relocated here
        import pipeline.attribution  # W5: hash-prune version constants relocated here
        cls._src = (
            pathlib.Path(
                inspect.getsourcefile(generate_weekly_pdfs)
            ).read_text(encoding='utf-8')
            + "\n"
            + pathlib.Path(
                inspect.getsourcefile(pipeline.fetch)
            ).read_text(encoding='utf-8')
            + "\n"
            + pathlib.Path(
                inspect.getsourcefile(pipeline.grouping)
            ).read_text(encoding='utf-8')
            + "\n"
            + pathlib.Path(
                inspect.getsourcefile(pipeline.cleanup)
            ).read_text(encoding='utf-8')
            + "\n"
            + pathlib.Path(
                inspect.getsourcefile(pipeline.attribution)
            ).read_text(encoding='utf-8')
        )

    def test_bug_a_rescue_gate_present_in_production(self):
        """Bug A rescue gate: is_subcontractor_sheet + kill switch."""
        # Look for the multiline conditional block — both names appear
        # consecutively gated by ``and`` in the production source.
        self.assertRegex(
            self._src,
            r'is_subcontractor_sheet\s*\n\s*and SUBCONTRACTOR_RATE_RECALC_PREACCEPTANCE_ENABLED',
            "Bug A rescue gate must be present in production",
        )
        self.assertIn('_subcontractor_rescue_price', self._src)

    def test_bug_b1_partitioning_gate_present_in_production(self):
        """Bug B1 partitioning gate.

        Phase 14 (D-14-06) extended this gate with a Helper #2 exclusion
        (``and not valid_helper2_row``) — the literal below was updated
        in the same change that extended the production gate.
        """
        self.assertIn(
            'if not is_subcontractor_row and not valid_helper_row '
            'and not valid_helper2_row:',
            self._src,
            "Bug B1 partitioning gate must be present in production",
        )

    def test_bug_b2_whitelist_signature_present_in_production(self):
        """Bug B2 whitelist kwarg signature."""
        self.assertRegex(
            self._src,
            r'variant_whitelist: set\[str\] \| None = None',
            "Bug B2 whitelist kwarg signature must be present in production",
        )

    def test_bug_c_reader_invocation_site_present_in_production(self):
        """Bug C reader invocation site (Phase 2: O(1) map read from bulk prefetch)."""
        # Phase 2 Plan 02: sub-helper attribution replaced by O(1) map read
        # from shared _attr_map (prefetch_attribution / D-03). The old
        # per-row lookup_attribution call is gone; resolve_claimer_sh
        # performs the map lookup via prefetched_map=_attr_map.
        #
        # Phase 2 Plan 05 (CR-01): the sub-helper call now routes the
        # prefetched map through the rpc_missing graceful-degradation gate
        # (``None if _attr_use_per_row_fallback else _attr_map``), so the
        # bare ``prefetched_map=_attr_map`` literal moved to the gated form.
        # The guard's intent — the sub-helper site reads the shared map —
        # is preserved by asserting the gated expression + the else-clause.
        self.assertIn('_resolve_claimer_sh', self._src)
        self.assertIn('_attr_use_per_row_fallback', self._src)
        self.assertRegex(
            self._src,
            r'prefetched_map=\(\s*\n?\s*None if _attr_use_per_row_fallback'
            r'\s*\n?\s*else _attr_map',
        )
        self.assertIn('SUBCONTRACTOR_HELPER_CLAIM_ATTRIBUTION_ENABLED', self._src)

    def test_hash_prune_version_constant_present_in_production(self):
        """Hash-prune version constant per D-17/D-19 (Plan 01.1-06 bumped to 2)."""
        self.assertRegex(
            self._src,
            r'(?m)^PHASE_1_1_HASH_PRUNE_VERSION = 2$',
            "Hash-prune version constant must be 2 in production (Plan 01.1-06)",
        )

    def test_hash_prune_helper_callable(self):
        """Helper function `_run_phase_1_1_hash_prune` is callable."""
        self.assertTrue(
            callable(generate_weekly_pdfs._run_phase_1_1_hash_prune),
            "Hash-prune helper must be a module-level callable",
        )

    def test_phase_1_1_hash_prune_pii_marker_present(self):
        """Prune-pass PII marker registered per [2026-05-15 12:00] rule 3."""
        self.assertIn(
            'Phase 1.1 hash-history prune',
            generate_weekly_pdfs._PII_LOG_MARKERS,
        )

    def test_helper2_shadow_partition_gate_present_in_production(self):
        """Phase 14 plan 14-06: Helper #2 sibling shadow partition gate
        (Pattern B — copied verbatim from the Helper #1 site)."""
        self.assertIn('_sub_is_valid_helper2_row', self._src)
        self.assertIn("('reduced_sub_helper2',", self._src)
        self.assertIn("'aep_billable_helper2'", self._src)


class TestSubcontractorHelper2ShadowRescue(unittest.TestCase):
    """Phase 14 plan 14-06 Task 1: the subcontractor Helper #2 shadow
    partition in ``group_source_rows``.

    Sibling of ``TestEndToEndPipeline``'s Helper #1 shadow tests above —
    same module-state setup, same synthetic-row idiom, extended with a
    Helper #2 row builder. Drives the production classifier (not a mirror
    class) per the Phase 1.1 Living Ledger entry rule (d).
    """

    _SUB_SHEET_ID = TestEndToEndPipeline._SUB_SHEET_ID
    _NON_SUB_SHEET_ID = TestEndToEndPipeline._NON_SUB_SHEET_ID

    setUp = TestEndToEndPipeline.setUp
    tearDown = TestEndToEndPipeline.tearDown
    _make_synth_helper_row = TestEndToEndPipeline._make_synth_helper_row
    _make_synth_non_helper_row = TestEndToEndPipeline._make_synth_non_helper_row

    def _make_synth_helper2_row(
        self,
        wr='19236776',
        helper2_foreman='ReplacementForeman2',
        units_price='$100.00',
        snapshot='2026-04-19',
        row_id=22345,
        source_sheet_id=None,
    ):
        """Synthetic Helper #2 row — sibling of
        ``TestEndToEndPipeline._make_synth_helper_row`` using the
        Helper #2 field names locked in plan 14-01
        (``__is_helper2_row`` / ``__helper2_foreman`` /
        ``__helper2_dept`` / ``__helper2_job``)."""
        return {
            '__row_id': row_id,
            'Work Request #': wr,
            'Weekly Reference Logged Date': '2026-04-19',
            'Snapshot Date': snapshot,
            'Units Completed?': True,
            'Foreman Helping? #2': helper2_foreman,
            'Helping Foreman #2 Completed Unit?': True,
            'Units Total Price': units_price,
            'CU': 'ANC-M',
            'Work Type': 'Inst',
            'Quantity': 2,
            '__effective_user': 'PrimaryForeman',
            '__assignment_method': 'FOREMAN_COLUMN',
            '__is_helper_row': False,
            '__helper_foreman': '',
            '__helper_dept': '',
            '__helper_job': '',
            '__is_helper2_row': True,
            '__helper2_foreman': helper2_foreman,
            '__helper2_dept': '600',
            '__helper2_job': 'JOB-B',
            '__is_vac_crew': False,
            '__source_sheet_id': source_sheet_id or self._SUB_SHEET_ID,
        }

    # ─── Core partition: shadow-only, no plain/primary leak ───────

    def test_subcontractor_helper2_row_emits_shadow_keys_only(self):
        """D-14-06 core assertion: a subcontractor Helper #2 completion
        emits exactly the two Helper #2 shadow keys — no plain ``helper2``
        key, and no subcontractor primary key (post-cutoff snapshot)."""
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value=None,  # no_history → falls back to current helper2
        ):
            row = self._make_synth_helper2_row(
                helper2_foreman='Drew_H2',
                snapshot='2026-04-19',  # post-AEP-cutoff
            )
            groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        self.assertTrue(
            any('REDUCEDSUB_HELPER2_Drew_H2' in k for k in keys),
            f"expected _REDUCEDSUB_HELPER2_ key; got: {keys}",
        )
        self.assertTrue(
            any('AEPBILLABLE_HELPER2_Drew_H2' in k for k in keys),
            f"expected _AEPBILLABLE_HELPER2_ key (post-cutoff); got: {keys}",
        )
        for k in keys:
            variant = groups[k][0].get('__variant', '')
            self.assertNotEqual(
                variant, 'helper2',
                f"no group should have variant='helper2' for a "
                f"subcontractor Helper #2 row; offending key={k!r}, "
                f"got: {keys}",
            )
            self.assertNotIn(
                variant, ('reduced_sub', 'aep_billable'),
                f"a Helper #2-completed subcontractor row must not "
                f"produce a primary variant group; offending key={k!r} "
                f"variant={variant!r}; got: {keys}",
            )
        self.assertEqual(
            len(keys), 2,
            f"expected exactly 2 shadow keys (reduced_sub_helper2 + "
            f"aep_billable_helper2); got: {keys}",
        )

    def test_subcontractor_helper2_row_pre_cutoff_emits_only_reducedsub_helper2(self):
        """Pre-cutoff snapshot — only ``_REDUCEDSUB_HELPER2_``, mirrors
        the Helper #1 pre-cutoff test."""
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value=None,
        ):
            row = self._make_synth_helper2_row(
                helper2_foreman='Drew_H2',
                snapshot='2026-04-11',  # pre-AEP-cutoff (cutoff is 04-12)
            )
            groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        self.assertTrue(
            any('REDUCEDSUB_HELPER2_Drew_H2' in k for k in keys),
            f"pre-cutoff: _REDUCEDSUB_HELPER2_ must be present; got: {keys}",
        )
        self.assertFalse(
            any('AEPBILLABLE_HELPER2_Drew_H2' in k for k in keys),
            f"pre-cutoff: _AEPBILLABLE_HELPER2_ must NOT be present; "
            f"got: {keys}",
        )

    def test_non_subcontractor_helper2_row_unaffected(self):
        """Non-subcontractor Helper #2 row is unchanged from plan 14-01:
        emits only the plain ``helper2`` key — no shadow variants (there
        is no subcontractor shadow leg to reach for a non-sub sheet)."""
        row = self._make_synth_helper2_row(
            source_sheet_id=self._NON_SUB_SHEET_ID
        )
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value={'helper2': 'FrozenHelper2'},
        ) as mock_lookup:
            groups = generate_weekly_pdfs.group_source_rows([row])
            mock_lookup.assert_not_called()
        keys = list(groups.keys())
        self.assertTrue(
            any('HELPER2_ReplacementForeman2' in k and 'REDUCEDSUB' not in k
                and 'AEPBILLABLE' not in k for k in keys),
            f"non-subcontractor Helper #2 row should emit the plain "
            f"_HELPER2_<name> key; got: {keys}",
        )

    # ─── VAC-crew precedence ───────────────────────────────────────

    def test_vac_crew_row_with_helper2_takes_vac_path(self):
        """D-14-06: the VAC-crew short-circuit wins over a Helper #2
        completion, exactly as it wins over Helper #1."""
        row = self._make_synth_helper2_row(
            helper2_foreman='Ignored_H2',
            snapshot='2026-04-19',
            row_id=55501,
        )
        row['__is_vac_crew'] = True
        row['__vac_crew_name'] = 'VacCrewMember'
        groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        self.assertTrue(
            any('VACCREW' in k for k in keys),
            f"VAC-crew row with Helper #2 data must take the VAC path; "
            f"got: {keys}",
        )
        for k in keys:
            variant = groups[k][0].get('__variant', '')
            self.assertNotIn(
                variant,
                ('helper2', 'reduced_sub_helper2', 'aep_billable_helper2'),
                f"VAC-crew short-circuit must win over Helper #2; "
                f"offending key={k!r} variant={variant!r}; got: {keys}",
            )

    # ─── Claimant resolution through the Helper #2 role ────────────

    def test_helper2_shadow_claimant_resolves_through_helper2_role(self):
        """D-14-06: the Helper #2 shadow claimant is resolved through
        ``resolve_claimer('helper2', ...)`` — the Helper #2 role added in
        plan 14-03 — not the primary or Helper #1 role."""
        with mock.patch(
            'billing_audit.writer.prefetch_attribution',
            return_value=({}, 'no_row'),
        ), mock.patch(
            'billing_audit.writer.resolve_claimer',
            return_value=ResolveOutcome(
                'use', 'FrozenHelper2', 'frozen', 'success'),
        ) as _rc:
            row = self._make_synth_helper2_row(helper2_foreman='CurrentHelper2')
            groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        self.assertTrue(
            any('REDUCEDSUB_HELPER2_FrozenHelper2' in k for k in keys),
            f"shadow file should use the frozen Helper #2 claimer; "
            f"got: {keys}",
        )
        helper2_role_calls = [
            c for c in _rc.call_args_list if c.args and c.args[0] == 'helper2'
        ]
        self.assertTrue(
            helper2_role_calls,
            "resolve_claimer must be called with the 'helper2' role for "
            "the Helper #2 shadow claimant",
        )

    # ─── O-14-A: helper2-wins conflict resolution ───────────────────

    def test_both_helper_slots_valid_helper2_wins(self):
        """O-14-A RESOLVED (.planning/phases/14-foreman-helper-2/
        14-DECISIONS.md, owner decision 2026-09-07): a subcontractor row
        with BOTH a valid Helper #1 AND a valid Helper #2 completion is
        resolved helper2-wins -- the Helper #2 shadow key is emitted and
        the Helper #1 shadow key is dropped for this row. Replaces this
        plan's placeholder 'takes no position' test now that Task 1's
        checkpoint has recorded the rule."""
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value=None,
        ):
            row = self._make_synth_helper_row(
                helper_foreman='Helper1Person',
                snapshot='2026-04-19',
            )
            row['__is_helper2_row'] = True
            row['__helper2_foreman'] = 'Helper2Person'
            row['__helper2_dept'] = '600'
            row['__helper2_job'] = 'JOB-B'
            groups = generate_weekly_pdfs.group_source_rows([row])
        keys = list(groups.keys())
        self.assertTrue(
            any('REDUCEDSUB_HELPER2_Helper2Person' in k for k in keys),
            f"Helper #2 shadow key must be emitted -- Helper #2 wins "
            f"per O-14-A; got: {keys}",
        )
        self.assertFalse(
            any(
                'REDUCEDSUB_HELPER_Helper1Person' in k and 'HELPER2' not in k
                for k in keys
            ),
            f"Helper #1 shadow key must be dropped for a conflicted "
            f"row per O-14-A; got: {keys}",
        )

    def test_both_helper_slots_valid_subcontractor_conflict_counted_once(self):
        """The subcontractor conflict is recorded exactly once — at the
        shadow leg — never double-counted against the plain-leg site,
        which defers subcontractor rows to the shadow leg (14-08 Task 2
        design: the plain leg's Helper #1/#2 blocks never emit a real
        key for a subcontractor row anyway)."""
        with mock.patch(
            'billing_audit.writer.lookup_attribution',
            return_value=None,
        ):
            row = self._make_synth_helper_row(
                helper_foreman='Helper1Person',
                snapshot='2026-04-19',
            )
            row['__is_helper2_row'] = True
            row['__helper2_foreman'] = 'Helper2Person'
            row['__helper2_dept'] = '600'
            row['__helper2_job'] = 'JOB-B'
            generate_weekly_pdfs.group_source_rows([row])
        self.assertEqual(_grouping.get_helper2_conflict_count(), 1)


class TestHelper2WrFilterAndExcludeWrs(unittest.TestCase):
    """Phase 14 code review CR-02: ``WR_FILTER`` (``_key_matches_wr``)
    and ``EXCLUDE_WRS`` (``_key_matches_excluded_wr``) in
    ``pipeline/grouping.py`` must recognize all three Helper #2
    group-key shapes -- plain ``_HELPER2_<name>``,
    ``_REDUCEDSUB_HELPER2_<name>``, and ``_AEPBILLABLE_HELPER2_<name>``
    -- exactly like the pre-Phase-14 Helper #1 clauses. Mirrors
    ``tests.test_primary_claim_attribution.TestWrFilterMatchesUserVariant``
    (Task 7's WR_FILTER/EXCLUDE_WRS regression pattern) using the
    Helper #2 fixtures from ``TestSubcontractorHelper2ShadowRescue``
    above.
    """

    _SUB_SHEET_ID = TestEndToEndPipeline._SUB_SHEET_ID
    _NON_SUB_SHEET_ID = TestEndToEndPipeline._NON_SUB_SHEET_ID

    _make_synth_helper2_row = (
        TestSubcontractorHelper2ShadowRescue._make_synth_helper2_row
    )

    def setUp(self):
        TestEndToEndPipeline.setUp(self)
        self._orig_wr_filter = list(generate_weekly_pdfs.WR_FILTER)
        self._orig_exclude_wrs = list(generate_weekly_pdfs.EXCLUDE_WRS)
        self._orig_test_mode = generate_weekly_pdfs.TEST_MODE

    def tearDown(self):
        generate_weekly_pdfs.WR_FILTER = self._orig_wr_filter
        generate_weekly_pdfs.EXCLUDE_WRS = self._orig_exclude_wrs
        generate_weekly_pdfs.TEST_MODE = self._orig_test_mode
        TestEndToEndPipeline.tearDown(self)

    def _rows(self):
        # WR under test: one subcontractor Helper #2 row (emits BOTH
        # _REDUCEDSUB_HELPER2_ and _AEPBILLABLE_HELPER2_, post-cutoff)
        # plus one non-subcontractor Helper #2 row (emits the plain
        # _HELPER2_ key). A second, untouched WR is the negative control.
        sub_row = self._make_synth_helper2_row(
            wr='19236776', helper2_foreman='Drew_H2',
            snapshot='2026-04-19', row_id=30001,
            source_sheet_id=self._SUB_SHEET_ID,
        )
        plain_row = self._make_synth_helper2_row(
            wr='19236776', helper2_foreman='Casey_H2',
            snapshot='2026-04-19', row_id=30002,
            source_sheet_id=self._NON_SUB_SHEET_ID,
        )
        other_wr_row = self._make_synth_helper2_row(
            wr='55555555', helper2_foreman='Other_H2',
            snapshot='2026-04-19', row_id=30003,
            source_sheet_id=self._NON_SUB_SHEET_ID,
        )
        return [sub_row, plain_row, other_wr_row]

    def test_wr_filter_retains_all_three_helper2_shapes(self):
        generate_weekly_pdfs.TEST_MODE = True
        generate_weekly_pdfs.WR_FILTER = ['19236776']
        with mock.patch(
            'billing_audit.writer.lookup_attribution', return_value=None,
        ):
            groups = generate_weekly_pdfs.group_source_rows(self._rows())
        keys = list(groups.keys())
        self.assertTrue(
            any('REDUCEDSUB_HELPER2_Drew_H2' in k for k in keys),
            f"WR_FILTER dropped the reduced_sub_helper2 group: {keys}",
        )
        self.assertTrue(
            any('AEPBILLABLE_HELPER2_Drew_H2' in k for k in keys),
            f"WR_FILTER dropped the aep_billable_helper2 group: {keys}",
        )
        self.assertTrue(
            any(
                'HELPER2_Casey_H2' in k and 'REDUCEDSUB' not in k
                and 'AEPBILLABLE' not in k
                for k in keys
            ),
            f"WR_FILTER dropped the plain helper2 group: {keys}",
        )
        self.assertFalse(
            any('Other_H2' in k for k in keys),
            f"WR_FILTER retained a group for an un-filtered WR: {keys}",
        )

    def test_exclude_wrs_drops_all_three_helper2_shapes(self):
        generate_weekly_pdfs.EXCLUDE_WRS = ['19236776']
        with mock.patch(
            'billing_audit.writer.lookup_attribution', return_value=None,
        ):
            groups = generate_weekly_pdfs.group_source_rows(self._rows())
        keys = list(groups.keys())
        self.assertFalse(
            any('Drew_H2' in k for k in keys),
            f"EXCLUDE_WRS failed to drop the subcontractor helper2 "
            f"shadow groups: {keys}",
        )
        self.assertFalse(
            any('Casey_H2' in k for k in keys),
            f"EXCLUDE_WRS failed to drop the plain helper2 group: {keys}",
        )
        self.assertTrue(
            any('Other_H2' in k for k in keys),
            f"EXCLUDE_WRS dropped a group for an un-excluded WR: {keys}",
        )


class TestHelper2ShadowExcelRendering(unittest.TestCase):
    """Phase 14 plan 14-06 Task 2: Helper #2 shadow workbook filename +
    REPORT DETAILS header rendering.

    Sibling of ``test_subcontractor_pricing.py``'s
    ``TestSubcontractorVariantFilenameSuffixes`` /
    ``TestSubcontractorHelperVariantDeptJobDisplay`` classes, added here
    per this plan's ``files_modified`` list (never touches
    ``test_subcontractor_pricing.py``).
    """

    def setUp(self):
        _ensure_smartsheet_mocked()
        self._tmpdir = tempfile.TemporaryDirectory()
        self._orig_output_folder = generate_weekly_pdfs.OUTPUT_FOLDER
        generate_weekly_pdfs.OUTPUT_FOLDER = self._tmpdir.name
        self._orig_rates = dict(generate_weekly_pdfs._SUBCONTRACTOR_RATES)
        generate_weekly_pdfs._SUBCONTRACTOR_RATES.clear()
        generate_weekly_pdfs._SUBCONTRACTOR_RATES['XYZ'] = {
            'cu_code': 'XYZ',
            'cu_wbs': '999',
            'compatible_unit_group': 'TestGroup',
            'reduced_install_price': 10.0,
            'reduced_remove_price': 5.0,
            'reduced_transfer_price': 2.5,
            'new_install_price': 20.0,
            'new_remove_price': 12.0,
            'new_transfer_price': 6.0,
        }

    def tearDown(self):
        generate_weekly_pdfs.OUTPUT_FOLDER = self._orig_output_folder
        self._tmpdir.cleanup()
        generate_weekly_pdfs._SUBCONTRACTOR_RATES.clear()
        generate_weekly_pdfs._SUBCONTRACTOR_RATES.update(self._orig_rates)

    def _make_group_row(self, variant, wr='10854049', week='2026-04-19',
                         snap='2026-04-19', helper2_foreman='', cu='XYZ',
                         work_type='Install', quantity=2, price='$0.00'):
        return {
            'Work Request #': wr,
            'Weekly Reference Logged Date': week,
            'Snapshot Date': snap,
            'Units Completed?': True,
            'Units Total Price': price,
            'CU': cu,
            'Work Type': work_type,
            'Quantity': quantity,
            'Customer Name': 'TestCustomer',
            'Foreman': 'TestForeman',
            'Dept #': '500',
            'Job #': 'J-1',
            '__effective_user': 'TestForeman',
            '__current_foreman': helper2_foreman or 'TestForeman',
            '__variant': variant,
            '__helper2_foreman': helper2_foreman,
            '__helper2_dept': '456' if helper2_foreman else '',
            '__helper2_job': 'J-9' if helper2_foreman else '',
            '__week_ending_date': datetime.datetime(2026, 4, 19),
        }

    def _read_detail(self, excel_path, label):
        """Return the REPORT DETAILS value (column G) for a given F-column label."""
        import openpyxl
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        for r in range(1, ws.max_row + 1):
            if ws.cell(row=r, column=6).value == label:  # column F = label
                return ws.cell(row=r, column=7).value      # column G = value
        return None

    # ─── Filename suffix ────────────────────────────────────────

    def test_aep_billable_helper2_filename_includes_sanitized_name(self):
        rows = [self._make_group_row('aep_billable_helper2', helper2_foreman='Jane Smith')]
        result = generate_weekly_pdfs.generate_excel(
            '041926_10854049_AEPBILLABLE_HELPER2_Jane_Smith', rows,
            datetime.datetime(2026, 4, 19), data_hash='deadbeefh2aep001',
        )
        filename = result[1]
        self.assertIn('_AEPBillable_Helper2_Jane_Smith_', filename)
        self.assertNotIn('_AEPBillable_Helper_Jane_Smith_', filename)

    def test_reduced_sub_helper2_filename_includes_sanitized_name(self):
        rows = [self._make_group_row('reduced_sub_helper2', helper2_foreman='Jane Smith')]
        result = generate_weekly_pdfs.generate_excel(
            '041926_10854049_REDUCEDSUB_HELPER2_Jane_Smith', rows,
            datetime.datetime(2026, 4, 19), data_hash='deadbeefh2rs001',
        )
        filename = result[1]
        self.assertIn('_ReducedSub_Helper2_Jane_Smith_', filename)
        self.assertNotIn('_ReducedSub_Helper_Jane_Smith_', filename)

    def test_aep_billable_helper2_raises_on_empty_helper2_foreman(self):
        """D-14-05 day-one defensive guard, mirroring the Helper #1
        shadow branch's raise-on-empty-foreman behavior."""
        rows = [self._make_group_row('aep_billable_helper2', helper2_foreman='')]
        with self.assertRaises(ValueError):
            generate_weekly_pdfs.generate_excel(
                '041926_10854049_AEPBILLABLE_HELPER2_', rows,
                datetime.datetime(2026, 4, 19), data_hash='deadbeefh2aep002',
            )

    def test_reduced_sub_helper2_raises_on_empty_helper2_foreman(self):
        rows = [self._make_group_row('reduced_sub_helper2', helper2_foreman='')]
        with self.assertRaises(ValueError):
            generate_weekly_pdfs.generate_excel(
                '041926_10854049_REDUCEDSUB_HELPER2_', rows,
                datetime.datetime(2026, 4, 19), data_hash='deadbeefh2rs002',
            )

    # ─── REPORT DETAILS header ──────────────────────────────────

    def test_reduced_sub_helper2_shows_helper2_dept_and_job(self):
        rows = [self._make_group_row('reduced_sub_helper2', helper2_foreman='Jane Smith')]
        result = generate_weekly_pdfs.generate_excel(
            '041926_10854049_REDUCEDSUB_HELPER2_Jane_Smith', rows,
            datetime.datetime(2026, 4, 19), data_hash='deadbeefh2rs003',
        )
        path = result[0]
        self.assertEqual(
            self._read_detail(path, 'Dept #:'), '456',
            "reduced_sub_helper2 file must show __helper2_dept (456), "
            "not primary Dept # (500)",
        )
        self.assertEqual(
            self._read_detail(path, 'Job #:'), 'J-9',
            "reduced_sub_helper2 file must show __helper2_job (J-9), "
            "not primary Job # (J-1)",
        )
        self.assertEqual(
            self._read_detail(path, 'Foreman:'), 'Jane Smith',
            "reduced_sub_helper2 file foreman must be the ATTRIBUTED "
            "claimer (current_foreman), preserving the same asymmetry "
            "as the Helper #1 shadow variants",
        )

    def test_aep_billable_helper2_shows_helper2_dept_and_job(self):
        rows = [self._make_group_row('aep_billable_helper2', helper2_foreman='Jane Smith')]
        result = generate_weekly_pdfs.generate_excel(
            '041926_10854049_AEPBILLABLE_HELPER2_Jane_Smith', rows,
            datetime.datetime(2026, 4, 19), data_hash='deadbeefh2aep003',
        )
        path = result[0]
        self.assertEqual(self._read_detail(path, 'Dept #:'), '456')
        self.assertEqual(self._read_detail(path, 'Job #:'), 'J-9')

    def test_plain_helper2_header_shows_raw_helper2_name(self):
        """Regression (plan 14-01): the plain 'helper2' branch shows the
        RAW __helper2_foreman, unlike the shadow variants above which
        show the attributed claimer."""
        rows = [self._make_group_row('helper2', helper2_foreman='Raw Helper2')]
        rows[0]['__current_foreman'] = 'ShouldNotAppear'
        result = generate_weekly_pdfs.generate_excel(
            '041926_10854049_HELPER2_Raw_Helper2', rows,
            datetime.datetime(2026, 4, 19), data_hash='deadbeefh2plain1',
        )
        path = result[0]
        self.assertEqual(self._read_detail(path, 'Foreman:'), 'Raw Helper2')


class TestHelper2FilenameParsingAndAggregatedHash(unittest.TestCase):
    """Phase 14 plan 14-06 Task 2: nested reserved-token filename parsing
    and multi-foreman aggregated-hash sub-bucketing for the Helper #2
    family, in ``pipeline/change_detection.py``.
    """

    def test_aep_billable_helper2_filename_parses_to_shadow_variant(self):
        result = generate_weekly_pdfs.build_group_identity(
            'WR_10854049_WeekEnding_041926_120000_AEPBillable_Helper2_'
            'Jane_Smith_abc123.xlsx'
        )
        self.assertIsNotNone(result)
        wr, week, variant, identifier = result
        self.assertEqual(variant, 'aep_billable_helper2')
        self.assertEqual(identifier, 'Jane_Smith')

    def test_reduced_sub_helper2_filename_parses_to_shadow_variant(self):
        result = generate_weekly_pdfs.build_group_identity(
            'WR_10854049_WeekEnding_041926_120000_ReducedSub_Helper2_'
            'Jane_Smith_abc123.xlsx'
        )
        self.assertIsNotNone(result)
        wr, week, variant, identifier = result
        self.assertEqual(variant, 'reduced_sub_helper2')
        self.assertEqual(identifier, 'Jane_Smith')

    def test_bare_aep_billable_user_filename_still_parses_unchanged(self):
        """Regression: a bare _AEPBillable_User_ filename (no Helper2
        token) must still parse to the bare aep_billable variant."""
        result = generate_weekly_pdfs.build_group_identity(
            'WR_10854049_WeekEnding_041926_120000_AEPBillable_User_'
            'ClaimerName_abc123.xlsx'
        )
        self.assertIsNotNone(result)
        wr, week, variant, identifier = result
        self.assertEqual(variant, 'aep_billable')
        self.assertEqual(identifier, 'ClaimerName')

    def test_bare_reduced_sub_helper_filename_still_parses_unchanged(self):
        """Regression: the Helper #1 shadow parse path (no Helper2 token
        present) must be completely unaffected by the new nested check."""
        result = generate_weekly_pdfs.build_group_identity(
            'WR_10854049_WeekEnding_041926_120000_ReducedSub_Helper_'
            'Jane_Smith_abc123.xlsx'
        )
        self.assertIsNotNone(result)
        wr, week, variant, identifier = result
        self.assertEqual(variant, 'reduced_sub_helper')
        self.assertEqual(identifier, 'Jane_Smith')

    def test_two_foreman_helper2_bucket_hash_changes_when_second_foreman_changes(self):
        """Pitfall 5 closure: an aggregation bucket holding rows from two
        distinct Helper #2 foremen must produce a hash that changes when
        the SECOND (non-first-sorted) foreman's identity changes — not
        only the first-sorted one."""
        from pipeline.change_detection import _compute_aggregated_content_hash

        def _row(foreman, price):
            return {
                '__variant': 'helper2',
                '__helper2_foreman': foreman,
                '__helper2_dept': '456',
                '__helper2_job': 'J-9',
                'Work Request #': '10854049',
                'Weekly Reference Logged Date': '2026-04-19',
                'Snapshot Date': '2026-04-19',
                'Units Completed?': True,
                'Units Total Price': price,
                'CU': 'XYZ',
                'Work Type': 'Install',
                'Quantity': 2,
                'Dept #': '500',
                'Job #': 'J-1',
            }

        baseline = [_row('Aaron_Helper2', '$100.00'), _row('Zack_Helper2', '$100.00')]
        # Only the SECOND-sorted foreman's row content changes.
        changed_second = [
            _row('Aaron_Helper2', '$100.00'),
            _row('Zack_Helper2', '$200.00'),
        ]
        hash_baseline = _compute_aggregated_content_hash(baseline)
        hash_changed = _compute_aggregated_content_hash(changed_second)
        self.assertNotEqual(
            hash_baseline, hash_changed,
            "a change to the second-sorted Helper #2 foreman's rows must "
            "change the aggregated hash — sub-bucketing must not depend "
            "on row sort order",
        )


class TestHelper2UploadRouting(unittest.TestCase):
    """Phase 14 plan 14-06 Task 3: the PPP dual-route gate for the
    ``reduced_sub_helper2`` variant in ``pipeline/upload.py``.

    Sibling of ``test_security_audit_followup.py``'s
    ``TestDualTargetSheetRouting`` -- added here per this plan's
    ``files_modified`` list (never touches
    ``test_security_audit_followup.py``).
    """

    @staticmethod
    def _make_kwargs(variant, *, wr_num='13792260', target_map=None,
                      target_map_ppp=None):
        """Minimal kwargs for ``_build_upload_tasks_for_group``.

        ``target_row`` is just a sentinel string for assertion
        readability -- the helper does not look inside it.
        """
        if target_map is None:
            target_map = {wr_num: f'row-TARGET-{wr_num}'}
        if target_map_ppp is None:
            target_map_ppp = {wr_num: f'row-PPP-{wr_num}'}
        return dict(
            variant=variant,
            wr_num=wr_num,
            target_map=target_map,
            target_map_ppp=target_map_ppp,
            excel_path=f'/tmp/{wr_num}.xlsx',
            filename=f'{wr_num}.xlsx',
            identifier='',
            file_identifier='',
            data_hash='abcdef0123456789',
            week_raw='041926',
            group_key=f'041926_{wr_num}_reduced_sub_helper2',
        )

    def test_reduced_sub_helper2_routes_to_both_sheets(self):
        """A reduced_sub Helper #2 group with a primary row present and
        a PPP target row available produces two upload tasks: one to
        the target sheet and one to the subcontractor PPP sheet."""
        kwargs = self._make_kwargs('reduced_sub_helper2')
        tasks = generate_weekly_pdfs._build_upload_tasks_for_group(**kwargs)
        self.assertEqual(len(tasks), 2)
        target_sheet_ids = {t['target_sheet_id'] for t in tasks}
        self.assertEqual(
            target_sheet_ids,
            {
                generate_weekly_pdfs.TARGET_SHEET_ID,
                generate_weekly_pdfs.SUBCONTRACTOR_PPP_SHEET_ID,
            },
        )

    def test_reduced_sub_helper1_still_routes_to_both_sheets(self):
        """Regression: a reduced_sub Helper #1 group still produces the
        same two tasks it produces today."""
        kwargs = self._make_kwargs('reduced_sub_helper')
        tasks = generate_weekly_pdfs._build_upload_tasks_for_group(**kwargs)
        self.assertEqual(len(tasks), 2)

    def test_aep_billable_helper2_routes_to_target_only(self):
        """An AEP-billable Helper #2 group produces one task, to the
        target sheet only, matching the AEP-billable Helper #1
        behavior."""
        kwargs = self._make_kwargs('aep_billable_helper2')
        tasks = generate_weekly_pdfs._build_upload_tasks_for_group(**kwargs)
        self.assertEqual(len(tasks), 1)
        self.assertEqual(
            tasks[0]['target_sheet_id'], generate_weekly_pdfs.TARGET_SHEET_ID,
        )

    def test_plain_helper2_routes_to_target_only(self):
        """A plain Helper #2 group produces one task, to the target
        sheet only."""
        kwargs = self._make_kwargs('helper2')
        tasks = generate_weekly_pdfs._build_upload_tasks_for_group(**kwargs)
        self.assertEqual(len(tasks), 1)
        self.assertEqual(
            tasks[0]['target_sheet_id'], generate_weekly_pdfs.TARGET_SHEET_ID,
        )

    def test_reduced_sub_helper2_without_ppp_target_row_skips_second_leg(self):
        """When no PPP target row exists for the work request, the
        second leg is skipped exactly as it is today, with no new
        failure mode."""
        kwargs = self._make_kwargs('reduced_sub_helper2', target_map_ppp={})
        tasks = generate_weekly_pdfs._build_upload_tasks_for_group(**kwargs)
        self.assertEqual(len(tasks), 1)
        self.assertEqual(
            tasks[0]['target_sheet_id'], generate_weekly_pdfs.TARGET_SHEET_ID,
        )


class TestSubcontractorWrScopeVariantGate(unittest.TestCase):
    """``_build_subcontractor_wr_scope`` gates on the authoritative
    ``__variant`` field (subcontractor variant set), not a ``'_REDUCEDSUB'``
    key substring — mirror of the Subproject D ``_build_primary_wr_scope``
    Codex-P1 consistency fix. A non-sub group whose claimer/helper NAME is an
    all-caps reserved token (``REDUCEDSUB`` / ``AEPBILLABLE``) must NOT enter
    the destructive subcontractor cleanup scope. Production rows always carry
    ``__variant`` (set at the ``group_source_rows`` emission site)."""

    def test_collects_all_subcontractor_variants(self):
        groups = {
            '041926_111_REDUCEDSUB': [
                {'Work Request #': '111', '__variant': 'reduced_sub'}],
            '041926_222_AEPBILLABLE': [
                {'Work Request #': '222', '__variant': 'aep_billable'}],
            '041926_333_REDUCEDSUB_HELPER_H': [
                {'Work Request #': '333', '__variant': 'reduced_sub_helper'}],
            '041926_444_AEPBILLABLE_HELPER_H': [
                {'Work Request #': '444', '__variant': 'aep_billable_helper'}],
            '041926_555_REDUCEDSUB_USER_C': [
                {'Work Request #': '555', '__variant': 'reduced_sub'}],
            '041926_666_REDUCEDSUB_HELPER2_H2': [
                {'Work Request #': '666', '__variant': 'reduced_sub_helper2'}],
            '041926_777_AEPBILLABLE_HELPER2_H2': [
                {'Work Request #': '777', '__variant': 'aep_billable_helper2'}],
        }
        scope = generate_weekly_pdfs._build_subcontractor_wr_scope(groups)
        self.assertEqual(
            scope, {'111', '222', '333', '444', '555', '666', '777'},
        )

    def test_helper2_only_wr_classified_subcontractor_active(self):
        """WR-01 (Phase 14 code review): a WR whose ONLY completed
        subcontractor rows this run are Helper #2 claims MUST be
        classified subcontractor-active -- pre-fix, ``__variant`` values
        of ``reduced_sub_helper2`` / ``aep_billable_helper2`` were absent
        from ``_SUBCONTRACTOR_SCOPE_VARIANTS`` and such a WR was silently
        dropped from the legacy off-contract / legacy-primary cleanup
        scope (``pipeline/orchestrate.py``'s ``_sub_scope`` gate)."""
        groups = {
            '041926_888_REDUCEDSUB_HELPER2_Only': [
                {'Work Request #': '888', '__variant': 'reduced_sub_helper2'}],
        }
        scope = generate_weekly_pdfs._build_subcontractor_wr_scope(groups)
        self.assertIn('888', scope)

    def test_excludes_non_subcontractor_variants(self):
        groups = {
            '041926_666_USER_Claimer': [
                {'Work Request #': '666', '__variant': 'primary'}],
            '041926_777_VACCREW_Vic': [
                {'Work Request #': '777', '__variant': 'vac_crew'}],
            '041926_888_HELPER_Help': [
                {'Work Request #': '888', '__variant': 'helper'}],
        }
        scope = generate_weekly_pdfs._build_subcontractor_wr_scope(groups)
        self.assertEqual(scope, set())

    def test_rejects_pathological_reducedsub_name(self):
        # Primary claimer literally named "REDUCEDSUB" → key contains the
        # _REDUCEDSUB substring, but __variant is 'primary'. The variant gate
        # must exclude it; the genuine sub group must still be collected.
        groups = {
            '041926_999_USER_REDUCEDSUB': [
                {'Work Request #': '999', '__variant': 'primary'}],
            '041926_123_REDUCEDSUB': [
                {'Work Request #': '123', '__variant': 'reduced_sub'}],
        }
        scope = generate_weekly_pdfs._build_subcontractor_wr_scope(groups)
        self.assertIn('123', scope)
        self.assertNotIn(
            '999', scope,
            "a non-sub group whose claimer name contains 'REDUCEDSUB' must "
            "NOT enter the subcontractor cleanup scope (variant gate)",
        )

    def test_empty_groups(self):
        self.assertEqual(
            generate_weekly_pdfs._build_subcontractor_wr_scope({}), set())


if __name__ == '__main__':
    unittest.main()
