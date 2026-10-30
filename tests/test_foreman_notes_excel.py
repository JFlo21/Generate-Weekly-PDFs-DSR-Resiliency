"""Phase 15 Plan 15-04: pilot Excel rendering (Foreman Notes + colour carry).

Pilot workbooks gain a ``Foreman Notes`` column I, A:I day headers, a red
TOTAL cell in I, fit-to-width printing and carried Smartsheet colours.  Every
switch-off / non-pilot workbook must stay cell-identical to the BASE golden
record (D-15-01, FN-12).  All data is synthetic (GF-1 style aliases only);
no network, no token.
"""

from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

_TESTS = Path(__file__).resolve().parent
_REPO_ROOT = _TESTS.parent
for _p in (str(_REPO_ROOT), str(_TESTS)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import openpyxl  # noqa: E402
from openpyxl.styles import numbers  # noqa: E402

import _foreman_notes_golden as golden  # noqa: E402
import generate_weekly_pdfs as gwp  # noqa: E402
from pipeline import foreman_notes as fn  # noqa: E402

RED = "00C00000"
YELLOW = "FEFF00"
STYLE_NAMES = [name for _col, name in fn.STYLE_COLUMNS]


def _fixture() -> dict:
    return json.loads(golden.FIXTURE_PATH.read_text(encoding="utf-8"))


def _groups() -> dict:
    return {name: (key, rows) for name, key, rows in golden.build_fixture_groups()}


def _pilot_state() -> fn.FnState:
    return fn.FnState(
        scope="pilot",
        allowlist=frozenset({"gf-1"}),
        pilot_wrs=frozenset({golden.WR}),
        gf_loaded=True,
    )


def _non_pilot_state() -> fn.FnState:
    return fn.FnState(
        scope="pilot",
        allowlist=frozenset({"gf-1"}),
        pilot_wrs=frozenset({"00000"}),
        gf_loaded=True,
    )


class _Base(unittest.TestCase):
    def setUp(self):
        fn.reset_state()
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.addCleanup(fn.reset_state)

    def render(self, name="primary", state=None, mutate=None):
        """Run the real ``generate_excel`` and return ``(path, dump)``."""
        key, rows = _groups()[name]
        rows = copy.deepcopy(rows)
        if mutate is not None:
            mutate(rows)
        fn.set_state(state) if state is not None else fn.reset_state()
        with golden.pinned_env(output_folder=self._tmp.name):
            path, *_rest = gwp.generate_excel(
                key, rows, golden.SNAPSHOT_DATE,
                data_hash=golden.hash_digests(rows)["data_hash"],
            )
        fn.reset_state()
        return path, golden.dump_workbook(path)

    @staticmethod
    def blocks(ws):
        """``[(header_row, first_cu_row, total_row)]`` per day block."""
        out = []
        for r in range(1, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == "Point Number":
                t = r + 1
                while ws.cell(row=t, column=1).value != "TOTAL":
                    t += 1
                out.append((r, r + 1, t))
        return out


class TracerTests(_Base):
    def test_pilot_primary_renders_note_and_carried_colour(self):
        style = {n: (YELLOW, None) for n in STYLE_NAMES if n != "N/A"}

        def mutate(rows):
            rows[0][fn.STYLE_KEY] = style

        path, _dump = self.render(state=_pilot_state(), mutate=mutate)
        ws = openpyxl.load_workbook(path).active
        blocks = self.blocks(ws)
        header_row, first_cu, _total = blocks[0]
        self.assertEqual(ws.cell(row=header_row, column=9).value, "Foreman Notes")
        self.assertEqual(
            ws.cell(row=first_cu, column=9).value,
            fn.norm_note(golden.NOTE_TEXTS[0]),
        )
        # Day header spans A:I. (The first day's header row overlaps the
        # summary merge G14:I14 and is skipped by safe_merge_cells exactly as
        # today, so the A:I merge is asserted on the second day.)
        second_header = blocks[1][0] - 1
        self.assertIn(f"A{second_header}:I{second_header}", ws.merged_cells)
        self.assertNotIn(
            f"A{second_header}:H{second_header}",
            [str(r) for r in ws.merged_cells.ranges])
        for col in (1, 2, 3, 4, 5, 6, 8):
            cell = ws.cell(row=first_cu, column=col)
            self.assertEqual(cell.fill.fill_type, "solid", col)
            self.assertEqual(cell.fill.fgColor.rgb[-6:], YELLOW, col)
        self.assertIsNone(ws.cell(row=first_cu, column=7).fill.fill_type)

        # second (unstyled) row equals the golden dump for A-H
        golden_cells = _fixture()["groups"]["primary"]["workbook"]["cells"]
        for col in "ABCDEFGH":
            coord = f"{col}{first_cu + 1}"
            self.assertEqual(
                golden.dump_workbook(path)["cells"].get(coord),
                golden_cells.get(coord), coord)

        # off again: identical to golden
        _p, off = self.render()
        self.assertEqual(off, _fixture()["groups"]["primary"]["workbook"])


if __name__ == "__main__":
    unittest.main()
