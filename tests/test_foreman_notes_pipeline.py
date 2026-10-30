"""Phase 15 Plan 15-05: pipeline wiring (run-start GF load, pre-render styles).

Two seams in ``pipeline/orchestrate.py`` call into ``pipeline/foreman_notes.py``:
``_fn.initialize(client)`` before PHASE 1 and ``_fn.attach_group_styles(...)``
just before ``generate_excel``.  Everything here is offline: a recording fake
client, fictional names (GF-style aliases) and fictional WR numbers only.
"""

from __future__ import annotations

import ast
import copy
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace as NS
from unittest import mock

_TESTS = Path(__file__).resolve().parent
_REPO_ROOT = _TESTS.parent
for _p in (str(_REPO_ROOT), str(_TESTS)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import openpyxl  # noqa: E402

import _foreman_notes_golden as golden  # noqa: E402
import generate_weekly_pdfs as gwp  # noqa: E402
from pipeline import change_detection as cd  # noqa: E402
from pipeline import config  # noqa: E402
from pipeline import foreman_notes as fn  # noqa: E402

ORCH_PATH = _REPO_ROOT / "pipeline" / "orchestrate.py"

SOURCE_SHEET = 7001
WR_COL_ID = 5180541294563204
GF_COL_ID = 2159630534528900
YELLOW = ",,,,,,,,,21,,,,,,,"
PILOT_WR = "90003"
OTHER_WR = "90004"

MAPPING = {
    "Pole #": 11, "CU": 12, "Work Type": 13, "CU Description": 14,
    "Unit of Measure": 15, "Quantity": 16, "Units Total Price": 17,
    "Foreman Notes": 18,
}


def _target_sheet():
    def cells(wr, gf):
        return NS(cells=[
            NS(column_id=WR_COL_ID, value=wr, display_value=wr),
            NS(column_id=GF_COL_ID, value=gf, display_value=gf),
        ])

    return NS(
        columns=[
            NS(id=WR_COL_ID, title="Work Request #"),
            NS(id=GF_COL_ID, title="General Foreman"),
        ],
        rows=[cells(PILOT_WR, "  avery   example "), cells(OTHER_WR, "NA")],
        total_row_count=2,
    )


class _Sheets:
    def __init__(self):
        self.calls = []

    def get_sheet(self, sheet_id, **kwargs):
        self.calls.append((sheet_id, kwargs))
        if sheet_id == config.TARGET_SHEET_ID:
            return _target_sheet()
        ids = [int(x) for x in kwargs["row_ids"].split(",")]
        return NS(rows=[
            NS(id=i, conditional_format=YELLOW, format_=None, cells=[])
            for i in ids
        ])


class _Client:
    def __init__(self):
        self.Sheets = _Sheets()


def _group(wr):
    """Golden-shaped primary rows re-labelled to a fictional WR."""
    rows = copy.deepcopy(golden._rows())
    for n, row in enumerate(rows, start=1):
        row["Work Request #"] = wr
        row["__source_sheet_id"] = SOURCE_SHEET
        row["__row_id"] = 1000 + n
    return f"{golden.WEEK_RAW}_{wr}", rows


class _Base(unittest.TestCase):
    def setUp(self):
        fn.reset_state()
        self.addCleanup(fn.reset_state)
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)

    def pilot_patches(self):
        stack = mock.patch.multiple(
            config, FOREMAN_NOTES_PILOT_GFS="Avery Example")
        scope = mock.patch.object(fn, "EFFECTIVE_SCOPE", "pilot")
        return stack, scope


class TracerTests(_Base):
    def test_pilot_path_end_to_end(self):
        client = _Client()
        patch_cfg, patch_scope = self.pilot_patches()
        with patch_cfg, patch_scope:
            state = fn.initialize(client)
        self.assertEqual(state.pilot_wrs, frozenset({PILOT_WR}))
        self.assertTrue(state.gf_loaded)

        maps = fn.build_sheet_column_maps(
            [{"id": SOURCE_SHEET, "column_mapping": MAPPING}])

        results = {}
        with golden.pinned_env(output_folder=self._tmp.name):
            for wr in (PILOT_WR, OTHER_WR):
                key, rows = _group(wr)
                hash_on = cd.calculate_data_hash(rows)
                fn.attach_group_styles(client, rows, maps)
                path, *_rest = gwp.generate_excel(
                    key, rows, golden.SNAPSHOT_DATE, data_hash=hash_on)
                # Hash of the same rows with the run state reset
                saved = fn.get_state()
                fn.reset_state()
                hash_off = cd.calculate_data_hash(rows)
                fn.set_state(saved)
                results[wr] = (path, hash_on, hash_off)

        # Exactly one target-sheet read and one style read.
        targets = [c for c in client.Sheets.calls
                   if c[0] == config.TARGET_SHEET_ID]
        styles = [c for c in client.Sheets.calls if c[0] == SOURCE_SHEET]
        self.assertEqual(len(targets), 1)
        self.assertEqual(len(styles), 1)
        self.assertEqual(styles[0][1].get("include"), "format")

        # Pilot workbook: note in I, carried fill on A-F/H/I of CU rows.
        ws = openpyxl.load_workbook(results[PILOT_WR][0]).active
        header = next(r for r in range(1, ws.max_row + 1)
                      if ws.cell(row=r, column=1).value == "Point Number")
        self.assertEqual(ws.cell(row=header, column=9).value, "Foreman Notes")
        first_cu = header + 1
        self.assertEqual(
            ws.cell(row=first_cu, column=9).value,
            fn.norm_note(golden.NOTE_TEXTS[0]))
        for col in (1, 2, 3, 4, 5, 6, 8, 9):
            cell = ws.cell(row=first_cu, column=col)
            self.assertEqual(cell.fill.fill_type, "solid", col)
            self.assertEqual(cell.fill.fgColor.rgb[-6:], "FEFF00", col)

        # Non-pilot workbook: no column I header, no carried fill.
        ws2 = openpyxl.load_workbook(results[OTHER_WR][0]).active
        for r in range(1, ws2.max_row + 1):
            self.assertNotEqual(ws2.cell(row=r, column=9).value, "Foreman Notes")
        header2 = next(r for r in range(1, ws2.max_row + 1)
                       if ws2.cell(row=r, column=1).value == "Point Number")
        for col in (1, 2, 3, 4, 5, 6, 8):
            fill = ws2.cell(row=header2 + 1, column=col).fill
            self.assertNotEqual(
                (fill.fgColor.rgb or "")[-6:] if fill.fill_type else None,
                "FEFF00", col)

        # Hash: differs for the pilot WR, identical for the other.
        self.assertNotEqual(results[PILOT_WR][1], results[PILOT_WR][2])
        self.assertEqual(results[OTHER_WR][1], results[OTHER_WR][2])


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
